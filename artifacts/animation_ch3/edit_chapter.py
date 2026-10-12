from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

root=Path(__file__).resolve().parent
src=Path(r'D:\งาน\project_บทที่3_ปรับปรุงพร้อมไดอะแกรม.docx')
out=root/'project_บทที่3_เพิ่มแอนิเมชันตัวละครและKiller.docx'
Image.open(root/'player_full.jpg').crop((503,132,1245,541)).save(root/'player_animator.png')
Image.open(root/'killer_full.jpg').crop((0,131,1040,570)).save(root/'killer_animator.png')
d=Document(src)
anchor=d.paragraphs[359]
added=[]
def p(text='',bold=False,center=False,page=False,size=16):
    x=anchor.insert_paragraph_before()
    x.paragraph_format.line_spacing=1
    x.paragraph_format.space_after=Pt(5)
    x.paragraph_format.page_break_before=page
    x.paragraph_format.keep_with_next=bold
    x.paragraph_format.widow_control=True
    x.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    if not bold and not center: x.paragraph_format.first_line_indent=Inches(.5)
    r=x.add_run(text); r.bold=bold; r.font.name='TH SarabunPSK'; r.font.size=Pt(size); r.font.color.rgb=RGBColor(0,0,0)
    rf=r._element.get_or_add_rPr().rFonts
    for key in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+key),'TH SarabunPSK')
    added.append(x)
    return x
def figure(file,caption):
    x=p(center=True); x.paragraph_format.keep_with_next=True
    x.add_run().add_picture(str(root/file),width=Inches(5.4 if file=='player_animator.png' else 6))
    p(caption,center=True)
def sub(text): return p(text,bold=True)

p('3.5.8 การพัฒนาระบบแอนิเมชันของตัวละครผู้เล่น',True,page=True,size=18)
p('ระบบแอนิเมชันใช้ Animator Controller ของ Unity เพื่อกำหนดสถานะภาพเคลื่อนไหวและเงื่อนไขเปลี่ยนสถานะ โดย Animation Clip เก็บลำดับภาพของตัวละคร ส่วน Animator เลือกคลิปที่ต้องแสดงตามพารามิเตอร์ซึ่งได้รับจากสคริปต์ วิธีนี้ทำให้การแสดงผลสัมพันธ์กับการเคลื่อนที่ และตรวจสอบเส้นทางการเปลี่ยนแอนิเมชันได้จากแผนผัง')
figure('player_animator.png','ภาพที่ 3-22 แผนผัง Animator ของตัวละครผู้เล่นจาก Unity')
p('จากภาพที่ 3-22 PlayerAnimator.controller มีสถานะ Idle และ Walk โดย Idle เป็นสถานะเริ่มต้น แสดงเป็นกล่องสีส้ม ลูกศรจาก Entry จึงเข้าสู่ Idle ก่อน ส่วน Walk แสดงภาพการเดิน ลูกศรระหว่างสองสถานะเป็น Transition ที่ใช้เงื่อนไข isWalking ชนิด Bool เพื่อเลือกแอนิเมชัน')
p('เมื่อ isWalking เป็น true ระบบเปลี่ยนจาก Idle ไป Walk และเมื่อเป็น false จะกลับจาก Walk ไป Idle ทั้งสอง Transition ปิด Has Exit Time และมี Transition Duration เท่ากับ 0 จึงเปลี่ยนสถานะได้ทันทีเมื่อเงื่อนไขเป็นจริง โดยไม่ต้องรอให้คลิปเดิมเล่นถึงจุดสิ้นสุด')
p('PlayerController ส่งค่าด้วย animator.SetBool(IsWalkingHash, IsMoving) จึงใช้สถานะการเคลื่อนที่ของผู้เล่นเป็นตัวควบคุมภาพจริง ใน Controller ปัจจุบันยังไม่มีสถานะ Run แยกต่างหาก การวิ่งจึงไม่ได้หมายความว่ามีคลิป Run ของผู้เล่นอยู่ในแผนผังนี้')
p('Any State เป็นจุดสำหรับกำหนดการเปลี่ยนจากสถานะใดก็ได้ แต่ของผู้เล่นยังไม่มี Transition ออกจากจุดนี้ ส่วน Entry เป็นจุดเริ่มต้น ไม่ใช่คลิปแอนิเมชัน การแยกความหมายของแต่ละกล่องช่วยให้ตีความแผนผังได้ถูกต้อง')

p('3.5.9 การพัฒนาระบบแอนิเมชันของ Killer',True,page=True,size=18)
p('Killer ใช้ KillerAnimator.controller เชื่อมกับ KillerAI เพื่อแสดงท่าทางตามพฤติกรรมของศัตรู ประกอบด้วย Idle สำหรับหยุดนิ่ง Walk สำหรับการเดิน Run สำหรับไล่ล่า Attack1 สำหรับโจมตี และ Scream สำหรับกรีดร้อง ภาพที่ 3-23 จับจากหน้าต่าง Animator ของ Unity โดยจัดตำแหน่งกล่องให้เห็นชื่อสถานะและลูกศรชัดเจน')
figure('killer_animator.png','ภาพที่ 3-23 แผนผัง Animator และ Parameters ของ Killer จาก Unity')
p('Entry เริ่มที่ Idle ส่วนการเดินใช้ isWalking ชนิด Bool และการวิ่งใช้ isRunning ชนิด Bool สคริปต์กำหนด isWalking เป็น true เมื่ออยู่ใน Patrol, Investigate, InvestigateSound, Search หรือ ReturnToPatrol และมีค่าสัมบูรณ์ของความเร็วแนวนอนมากกว่า 0.01 ส่วน isRunning เป็น true เมื่ออยู่ใน Chase หรือ Attack')
p('การโจมตีใช้ attack1 ชนิด Trigger โดย KillerAI เรียก SetTrigger เมื่อโจมตีตามรอบคูลดาวน์ หรือเริ่มตรวจค้นจุดซ่อนที่สงสัย จึงเรียกท่า Attack1 เป็นเหตุการณ์แต่ละครั้ง ขณะที่ scream เป็น Trigger ซึ่งเรียกเมื่อระบบสั่งกรีดร้องหรือใช้ความสามารถหยุดผู้เล่นตามเงื่อนไขของศัตรู')
p('Any State เชื่อมไป Scream เพื่อให้การกรีดร้องเริ่มได้จากสถานะใดก็ได้ ส่วน Attack1 รับ Transition จาก Idle, Walk และ Run จึงไม่ได้ใช้ Any State สำหรับการโจมตี เมื่อ Attack1 หรือ Scream เล่นครบตาม Exit Time เท่ากับ 1 จะกลับ Idle แล้วตรวจพารามิเตอร์การเคลื่อนที่เพื่อเลือกสถานะถัดไป')

p('3.5.10 เงื่อนไขการเปลี่ยนแอนิเมชันและการเชื่อมกับ FSM',True,page=True,size=18)
p('ลูกศรในภาพที่ 3-23 แสดงทิศทางที่ Animator เปลี่ยนสถานะได้ เงื่อนไขของ Transition ต้องสอดคล้องกับค่าที่ KillerAI ส่งเข้ามา โดย Controller ปัจจุบันกำหนดเส้นทางดังต่อไปนี้')
for t in [
'Idle ไป Walk เมื่อ isWalking เป็น true และ Walk กลับ Idle เมื่อ isWalking เป็น false',
'Idle หรือ Walk ไป Run เมื่อ isRunning เป็น true และ Run กลับ Idle เมื่อ isRunning เป็น false',
'Run ไป Walk เมื่อ isRunning เป็น false และ isWalking เป็น true โดยเงื่อนไขทั้งสองต้องเป็นจริงร่วมกัน',
'Idle, Walk หรือ Run ไป Attack1 เมื่อได้รับ Trigger attack1 และ Attack1 กลับ Idle ด้วย Has Exit Time ที่ค่า 1',
'Any State ไป Scream เมื่อได้รับ Trigger scream และ Scream กลับ Idle ด้วย Has Exit Time ที่ค่า 1']:
    p(t)
p('Transition ระหว่างการเคลื่อนที่ของ Killer มี Duration เท่ากับ 0.1 วินาที และปิด Has Exit Time เพื่อให้ตอบสนองต่อพฤติกรรมได้โดยไม่รอจบรอบคลิป การเข้าสู่ Attack1 ใช้ Duration 0.05 วินาที ส่วน Any State ไป Scream ใช้ Duration 0 เพื่อเริ่มกรีดร้องทันที การกลับจาก Attack1 และ Scream ใช้ Duration 0.1 วินาที')
p('FSM ใน KillerAI รับหน้าที่ตัดสินใจ เช่น ลาดตระเวน ค้นหา ไล่ล่า หรือโจมตี ส่วน Animator รับหน้าที่แสดงภาพของการตัดสินใจนั้น จึงไม่จำเป็นต้องมีจำนวนสถานะเท่ากัน ตัวอย่างเช่น Patrol และ Search ใช้ภาพ Walk ร่วมกันได้ แต่ยังมีตรรกะเลือกเป้าหมายต่างกันใน FSM')
p('สำหรับผู้เล่น หลักการเดียวกันใช้กับ Player FSM ซึ่งควบคุมการรับคำสั่งและข้อจำกัดในสถานะ Normal, Hidden, Frozen, Hurt, UsingItem, Transitioning และ Dead ส่วน PlayerAnimator ปัจจุบันเลือกภาพระหว่าง Idle กับ Walk ตาม IsMoving การเพิ่มสถานะใน FSM จึงไม่ทำให้เกิดคลิปแอนิเมชันใหม่โดยอัตโนมัติ')
p('การคำนวณความเสียหายอยู่ใน KillerAI ซึ่งเรียก playerHealth.TakeDamage ก่อนส่ง Trigger attack1 แอนิเมชันโจมตีจึงเป็นภาพประกอบเหตุการณ์ ไม่ใช่ตัวกำหนดความเสียหายด้วย Animation Event ในส่วนโค้ดนี้ หากต้องการให้ความเสียหายตรงกับเฟรมที่อาวุธสัมผัสผู้เล่น จึงค่อยออกแบบ Animation Event เพิ่มเติม')
p('เมื่อ isRunning เป็น false และ isWalking เป็น true เส้นจาก Run ไป Idle กับเส้นไป Walk อาจผ่านเงื่อนไขพร้อมกัน จึงควรตรวจลำดับ Transition เพื่อให้ผลการเปลี่ยนภาพสอดคล้องกับพฤติกรรมที่ต้องการ')

p('3.5.11 เหตุผลในการเลือกชนิดพารามิเตอร์ของ Animator',True,page=True,size=18)
p('Unity มีพารามิเตอร์ Float, Int, Bool และ Trigger เพื่อแทนข้อมูลต่างลักษณะ การเลือกชนิดขึ้นอยู่กับว่าเป็นค่าต่อเนื่อง รหัสสถานะ เงื่อนไขที่คงอยู่ หรือเหตุการณ์เฉพาะครั้ง ไม่จำเป็นต้องใช้ครบทั้งสี่ชนิดใน Controller เดียว โครงการนี้ใช้ Bool และ Trigger จริง ส่วน Float และ Int อธิบายเป็นทางเลือกสำหรับการขยายระบบ')
sub('Float สำหรับค่าต่อเนื่อง')
p('Float เก็บตัวเลขที่มีทศนิยม เหมาะกับค่าความเร็ว เช่น 0, 1.5 และ 3.2 หากเพิ่มพารามิเตอร์ Speed จะส่งค่าด้วย SetFloat และใช้เปรียบเทียบกับเกณฑ์หรือควบคุม Blend Tree เพื่อผสมท่ายืน เดิน และวิ่งตามความเร็วได้ เหตุผลที่เลือก Float คือรักษารายละเอียดระหว่างค่าความเร็วซึ่ง Int จะเก็บไม่ได้ ปัจจุบันยังไม่มีพารามิเตอร์ Float ใน Animator ทั้งสองตัว')
sub('Int สำหรับรหัสหรือทางเลือกหลายค่า')
p('Int เก็บจำนวนเต็ม เหมาะกับการแทนโหมดที่แยกกัน เช่น MovementMode เท่ากับ 0 สำหรับยืน 1 สำหรับเดิน และ 2 สำหรับวิ่ง แล้วส่งค่าด้วย SetInteger เหตุผลที่เลือก Int คือใช้ค่าเดียวระบุทางเลือกหลายแบบ ลดการใช้ Bool หลายตัวที่อาจเป็น true พร้อมกัน อย่างไรก็ตาม ตัวแปร int ที่เก็บผล Animator.StringToHash เป็นเพียงรหัสอ้างอิงชื่อพารามิเตอร์ ไม่ได้ทำให้พารามิเตอร์นั้นเป็นชนิด Int และ Controller ปัจจุบันยังไม่ได้ใช้ Int')
sub('Bool สำหรับเงื่อนไขที่ต้องคงสถานะ')
p('Bool มีค่า true หรือ false เหมาะกับคำถามว่ากำลังเดินหรือกำลังวิ่งหรือไม่ จึงใช้ isWalking ในผู้เล่น และ isWalking กับ isRunning ใน Killer โดยส่งค่าด้วย SetBool เหตุผลที่เลือก Bool คือระบบต้องรักษาค่าขณะพฤติกรรมยังดำเนินอยู่ และเปลี่ยนกลับเมื่อพฤติกรรมหยุด ทำให้การเข้าและออกจาก Walk หรือ Run ตรวจสอบได้ตรงไปตรงมา')
sub('Trigger สำหรับเหตุการณ์เฉพาะครั้ง')
p('Trigger เหมาะกับคำสั่งเริ่มการกระทำ เช่น โจมตีหรือกรีดร้อง จึงใช้ attack1 และ scream ผ่าน SetTrigger เมื่อ Transition ใช้ Trigger แล้ว Controller จะรีเซ็ตค่าโดยอัตโนมัติ หากยังไม่มี Transition ใช้ค่านั้น อาจต้องเรียก ResetTrigger เพื่อยกเลิกคำสั่งที่ค้างอยู่ เหตุผลที่เลือก Trigger คือไม่ต้องเปิด Bool ค้างไว้และคอยสั่งปิดเองทุกครั้ง ส่วนการออกจากคลิปยังต้องกำหนด Transition และ Exit Time ให้เหมาะสม')
p('แหล่งตรวจสอบระบบคือ PlayerAnimator.controller, KillerAnimator.controller, PlayerController.cs และ KillerAI.cs ในโปรเจกต์ SideScrollingEscape ส่วนชนิดพารามิเตอร์อ้างอิง Unity Manual เรื่อง Animation Parameters ที่ https://docs.unity3d.com/6000.4/Documentation/Manual/AnimationParameters.html ตรวจสอบวันที่ 4 ตุลาคม 2569',size=14)

# Keep the original following section at a clear page boundary.
anchor.paragraph_format.page_break_before=True
d.save(out)
# Small proof document permits separate verification of inserted pages.
proof=Document(src)
body=proof._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'): body.remove(child)
for x in added: body.insert(len(body)-1,deepcopy(x._p))
# Copy updated package so images and their relationship ids remain valid.
import zipfile
from lxml import etree
with zipfile.ZipFile(out) as zin, zipfile.ZipFile(root/'added_pages_proof.docx','w',zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data=etree.tostring(proof._element,xml_declaration=True,encoding='UTF-8',standalone=True) if item.filename=='word/document.xml' else zin.read(item.filename)
        zout.writestr(item,data)
print(out)

