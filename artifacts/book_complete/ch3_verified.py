from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree
r=Path(__file__).resolve().parent
d=Document(r/'project_บท4_ผลทดสอบจริง_PlayMode.docx')
b=list(d.element.body);txt=lambda e:''.join(e.xpath('.//w:t/text()'))
s=next(i for i,e in enumerate(b) if txt(e)=='บทที่ 3');z=next(i for i,e in enumerate(b) if i>s and txt(e).startswith('บทที่ 4 ผล'))
before=[etree.tostring(e) for e in b[:s]];after=[etree.tostring(e) for e in b[z:]]
results={}
for t in d.tables:
 if t._tbl in b[z:]:
  for row in t.rows[1:]:
   if len(row.cells)==5 and row.cells[0].text.startswith('TC-'):
    results[row.cells[0].text]=(row.cells[3].text,row.cells[4].text)
count=0
for t in d.tables:
 if t._tbl not in b[s:z]:continue
 if t.rows[0].cells[-1].text!='ผลที่บันทึกและสถานะ':continue
 for row in t.rows[1:]:
  id=row.cells[0].text
  result,status=results[id]
  c=row.cells[-1];c.text=result+'\nสถานะ: '+status
  for p in c.paragraphs:
   for run in p.runs:
    run.font.name='TH SarabunPSK';run.font.size=Pt(14)
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+key),'TH SarabunPSK')
  count+=1
for p in d.paragraphs:
 if p._p not in b[s:z]:continue
 if p.text.startswith('ตารางการทดสอบบันทึกผลจากการตรวจโค้ด'):
  p.text='ตารางต่อไปบันทึกผลการทดสอบใน Unity 6000.4.9f1 เมื่อวันที่ 4 ตุลาคม 2569 โดยใช้ผลชุดเดียวกับหัวข้อ 4.10 ทดสอบใน Play Mode ด้วยสคริปต์ชั่วคราวควบคุมอินพุต ตำแหน่ง เวลา และเงื่อนไขขององค์ประกอบ สถานะผ่านหมายถึงผ่านกรณีควบคุมที่ระบุ รวม 30 กรณีของระบบ ไม่ใช่การรับรองการเล่นครบทุกเส้นทางหรือผลประเมินจากผู้ใช้งานจริง'
  for run in p.runs:run.font.name='TH SarabunPSK';run.font.size=Pt(16)
out=r/'project_บท3และ4_ผลทดสอบล่าสุด.docx';d.save(out)
c=Document(out);bb=list(c.element.body);ss=next(i for i,e in enumerate(bb) if txt(e)=='บทที่ 3');zz=next(i for i,e in enumerate(bb) if i>ss and txt(e).startswith('บทที่ 4 ผล'))
assert before==[etree.tostring(e) for e in bb[:ss]]
assert after==[etree.tostring(e) for e in bb[zz:]]
assert count==30
print('Filled 30 chapter 3 cases from verified chapter 4 results; other chapters unchanged.')
