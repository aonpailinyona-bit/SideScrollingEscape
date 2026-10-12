from pathlib import Path
from copy import deepcopy
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
import re

r=Path(__file__).resolve().parent
base=Document(next((r.parent/'animation_ch3').glob('*3และ4*.docx')))
donor=Document(r/'project_ฉบับสมบูรณ์พร้อมภาพ.docx')
def txt(e):return ''.join(e.xpath('.//w:t/text()'))
def bounds(d):
    b=list(d.element.body)
    start=next(i for i,e in enumerate(b) if txt(e) in ['บทที่ 4','บทที่ 4 ผลการดำเนินงาน'])
    end=next(i for i,e in enumerate(b) if i>start and txt(e) in ['บทที่ 5','บทที่ 5 สรุปผลโครงงานและข้อเสนอแนะ'])
    return b,start,end
old,s,e=bounds(base);new,ds,de=bounds(donor)
before=[etree.tostring(x) for x in old[:s]]
after=[etree.tostring(x) for x in old[e:]]
rels={}
def clone(node):
    x=deepcopy(node)
    for el in x.iter():
        for attr in [qn('r:embed'),qn('r:link'),qn('r:id')]:
            rid=el.get(attr)
            if not rid or rid not in donor.part.rels:continue
            if rid not in rels:
                rel=donor.part.rels[rid]
                if rel.reltype==RT.IMAGE:
                    part=base.part.package.get_or_add_image_part(BytesIO(rel.target_part.blob))
                    rels[rid]=base.part.relate_to(part,RT.IMAGE)
                elif rel.is_external:rels[rid]=base.part.relate_to(rel.target_ref,rel.reltype,True)
                else:raise RuntimeError('Unexpected relationship '+rel.reltype)
            el.set(attr,rels[rid])
    # Use local formatting, preserving every existing style outside chapter 4.
    for st in x.xpath('.//w:pStyle'):
        style=st.get(qn('w:val'))
        if style.startswith('Heading') or style in ['FigureCaption','TableCaption']:
            pr=st.getparent();pr.remove(st)
            if style.startswith('Heading'):
                keep=etree.Element(qn('w:keepNext'));pr.append(keep)
    for el in x.xpath('.//w:t'):
        if el.text:el.text=el.text.replace('ภาคผนวก ข','หัวข้อ 4.11.2')
    heading=bool(re.match(r'^(บทที่ 4|4\.\d+(?:\.\d+)* )',txt(x)))
    size=14 if x.tag==qn('w:tbl') else (18 if heading else 16)
    for run in x.xpath('.//w:r'):
        pr=run.find(qn('w:rPr'))
        if pr is None:pr=etree.Element(qn('w:rPr'));run.insert(0,pr)
        for tag in ['rFonts','sz','szCs','color','rStyle']:
            for oldpr in pr.findall(qn('w:'+tag)):pr.remove(oldpr)
        rf=etree.SubElement(pr,qn('w:rFonts'))
        for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'TH SarabunPSK')
        for tag in ['sz','szCs']:
            sz=etree.SubElement(pr,qn('w:'+tag));sz.set(qn('w:val'),str(size*2))
        color=etree.SubElement(pr,qn('w:color'));color.set(qn('w:val'),'000000')
    return x
content=[clone(x) for x in new[ds:de]]
# Include the requested mock scores within chapter 4, not elsewhere in the book.
bs=next(i for i,x in enumerate(new) if txt(x).startswith('ข.1 ข้อมูลสมมติ'))
be=next(i for i,x in enumerate(new) if i>bs and txt(x).startswith('ข.2 แบบประเมิน'))
mock=[clone(x) for x in new[bs:be]]
for node in mock:
    for t in node.xpath('.//w:t'):
        if t.text:t.text=t.text.replace('ข.1 ข้อมูลสมมติสำหรับแสดงวิธีคำนวณ','4.11.2 ตัวอย่างคะแนนสมมติสำหรับแสดงวิธีคำนวณ').replace('ตารางที่ ข-1','ตารางที่ 4-6').replace('ตารางที่ ข-2','ตารางที่ 4-7')
idx=next(i for i,x in enumerate(content) if txt(x).startswith('4.12 สรุปผล'))
content[idx:idx]=mock
anchor=old[e]
for node in old[s:e]:base.element.body.remove(node)
for node in content:anchor.addprevious(node)
sep=etree.Element(qn('w:p'));run=etree.SubElement(sep,qn('w:r'));br=etree.SubElement(run,qn('w:br'));br.set(qn('w:type'),'page');anchor.addprevious(sep)
# Start the updated chapter on a fresh page, and preserve the original chapter 5 boundary.
pr=content[0].find(qn('w:pPr'))
br=etree.Element(qn('w:pageBreakBefore'));pr.append(br)
out=r/'project_แก้เฉพาะบทที่4.docx'
base.save(out)
check=Document(out);bb,ss,ee=bounds(check)
assert before==[etree.tostring(x) for x in bb[:ss]],'Content before chapter 4 changed'
assert after==[etree.tostring(x) for x in bb[ee:]],'Content after chapter 4 changed'
assert not any('[แทรกภาพ' in txt(x) or '[กรอก' in txt(x) for x in bb[ss:ee])
assert sum(len(x.xpath('.//w:drawing')) for x in bb[ss:ee])==21
print('Chapter 4 updated; all other body elements preserved exactly')
