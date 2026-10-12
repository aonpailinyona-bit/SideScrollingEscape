from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree
r=Path(__file__).resolve().parent
evidence={x['id']:x for x in json.loads((r/'verification.json').read_text(encoding='utf-8-sig'))['rows']}
d=Document(r/'project_บท3และ4_กรอกตารางผล.docx')
b=list(d.element.body);txt=lambda e:''.join(e.xpath('.//w:t/text()'))
s=next(i for i,e in enumerate(b) if txt(e).startswith('บทที่ 4 ผล'))
z=next(i for i,e in enumerate(b) if i>s and txt(e)=='บทที่ 5')
before=[etree.tostring(e) for e in b[:s]];after=[etree.tostring(e) for e in b[z:]]
descriptions={
'TC-01':'กด D ได้ความเร็ว 4.5 และหันขวา กด A ได้ −4.5 และหันซ้าย',
'TC-02':'ความเร็วเดิน 4.5 เพิ่มเป็น 5.7 ขณะวิ่ง ค่า Stamina ลด และส่งเสียงประเภท Sprint',
'TC-03':'กระโดดได้แรงยกครั้งแรก การสั่งกระโดดซ้ำทันทีขณะลอยไม่เพิ่มแรงยก',
'TC-04':'ขึ้นบันไดด้วยความเร็ว 4 และลงด้วย −4 กระโดดออกได้และคืนค่าแรงโน้มถ่วง',
'TC-05':'เปิดปิดได้ แบตเตอรี่ลดขณะเปิด เมื่อเหลือศูนย์ไฟดับและเปิดซ้ำไม่ได้',
'TC-FSM-P01':'ToggleHide เปลี่ยนเป็น Hidden และล็อกการเคลื่อนไหว',
'TC-FSM-P02':'ToggleHide ออกจากที่ซ่อน กลับ Normal และคืนการควบคุม',
'TC-FSM-P03':'รับความเสียหาย 10 ทำให้ HP 100 เป็น 90 เข้า Hurt และกลับ Normal หลังรอ 2.2 วินาที',
'TC-FSM-P04':'Freeze เปลี่ยนเป็น Frozen กด X ครบ 2 ครั้งตามค่าทดสอบแล้วกลับ Normal',
'TC-FSM-P05':'เปิดสถานะคลังเป็น UsingItem และล็อกการเคลื่อนไหว ปิดแล้วกลับ Normal',
'TC-FSM-P06':'BeginTransition ล็อกคำสั่ง และเมื่อประตูสำเร็จผู้เล่นเข้าสู่ Transitioning',
'TC-FSM-P07':'รับความเสียหายจน HP หมด ผู้เล่นเป็น Dead และสถานะเกมเป็น Dead',
'TC-FSM-E01':'ส่งเสียง Rock ในระยะได้ยิน ทำให้ Patrol เปลี่ยนเป็น Listening',
'TC-FSM-E02':'หมดเวลาฟังแล้วเข้า InvestigateSound และมีความเร็วเคลื่อนที่ไปยังเสียง',
'TC-FSM-E03':'เมื่อตำแหน่งถึงต้นกำเนิดเสียง เปลี่ยนเป็น ConfirmSound',
'TC-FSM-E04':'ผู้เล่นอยู่ 2 หน่วยด้านที่ Killer หันและไม่มีสิ่งกีดขวาง ตรวจพบจริงแล้วเข้า Chase',
'TC-FSM-E05':'ผู้เล่นอยู่ในระยะโจมตี Chase เข้า Attack จากนั้น HP 100 ลดเป็น 75 และเข้า AttackRecover',
'TC-FSM-E06':'คง AttackRecover ก่อนเวลาสิ้นสุด และกลับ Chase เมื่อครบเวลาขณะที่ยังตรวจพบเป้าหมาย',
'TC-FSM-E07':'สูญเสียเป้าหมายหลังช่วงความจำแล้วเข้า Investigate เมื่อหมดเวลาตรวจสอบเข้า Search',
'TC-FSM-E08':'หมดเวลาค้นหาแล้วเปลี่ยนจาก Search เป็น ReturnToPatrol',
'TC-FSM-E09':'ส่ง Rock ซ้ำจุดเดิม 3 ครั้ง เมื่อยืนยันเสียงแล้วเข้า Confused',
'TC-FSM-E10':'เรียกเหตุการณ์กรีดร้องแล้ว Killer เข้า Scream และผู้เล่นถูก Stun จนควบคุมไม่ได้ชั่วคราว',
'TC-06':'Collect เพิ่มกุญแจจาก 0 เป็น 1 และเพิ่มรายการใน Inventory',
'TC-07':'ขณะมีกุญแจ 1 จาก 5 ประตูยัง Locked และแจ้งว่าขาดอีก 4 ดอก',
'TC-08':'กด E ค้างตามระยะเวลาปลดล็อก ประตูเป็น Completed และเกมเป็น LevelComplete',
'TC-09':'ยกเลิกและคืนความคืบหน้าเป็นศูนย์เมื่อปล่อย E ออกจากระยะ รับความเสียหาย หรือ Killer อยู่ใกล้',
'TC-10':'กับดักทำงาน ปิด Collider ส่ง Glass จน Killer เข้า Listening และ AudioSource เล่นเสียงที่มีข้อมูลเสียงจริง',
'TC-11':'สุ่มตำแหน่ง 5 รอบ ตำแหน่งกุญแจเปลี่ยน และ Raycast พบพื้น Ground ใต้กุญแจ',
'TC-12':'ความเข้มเตือนเพิ่มจาก 0 เป็นประมาณ 0.929 เมื่อเข้าใกล้ ขอบจอมีค่า Alpha และคลิปหัวใจมีข้อมูลเสียง',
'TC-13':'แจ้งผ่านด่านสุดท้ายและสั่งดำเนินต่อแล้วสถานะเกมเป็น GameWon',
'AN-P01':'isWalking false → true → false ทำให้ Idle → Walk → Idle',
'AN-K01':'isWalking true แสดง Walk และ isRunning true เปลี่ยนเป็น Run',
'AN-K02':'Trigger attack1 เข้า Attack1 และกลับ Idle หลังคลิป',
'AN-K03':'Trigger scream จาก Idle, Walk, Run และ Attack1 เข้า Scream และกลับ Idle ได้ทุกกรณี',
'AN-K04':'ขณะ Run และ Bool ทั้งสองเป็น true เมื่อ isRunning เป็น false ภาพกลับ Walk'
}
def style_cell(c):
 for p in c.paragraphs:
  p.paragraph_format.keep_with_next=False
  for run in p.runs:
   run.font.name='TH SarabunPSK';run.font.size=Pt(14)
   rf=run._element.get_or_add_rPr().get_or_add_rFonts()
   for key in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+key),'TH SarabunPSK')
for t in d.tables:
 if t._tbl not in b[s:z]:continue
 for row in t.rows[1:]:
  id=row.cells[0].text.strip()
  if id not in descriptions:continue
  result=evidence[id]
  if len(row.cells)==5:
   row.cells[3].text=descriptions[id];row.cells[4].text='ผ่าน' if result['status']=='PASS' else 'ไม่ผ่าน'
   style_cell(row.cells[3]);style_cell(row.cells[4])
  elif len(row.cells)==3:
   row.cells[2].text=('ผ่าน: ' if result['status']=='PASS' else 'ไม่ผ่าน: ')+descriptions[id];style_cell(row.cells[2])
scope='ทดสอบวันที่ 4 ตุลาคม 2569 ด้วย Unity 6000.4.9f1 ใน Play Mode โดยใช้สคริปต์ชั่วคราวควบคุมตำแหน่ง อินพุต เวลา และเงื่อนไขขององค์ประกอบ แล้วตรวจสถานะ ค่าความเร็ว HP ทรัพยากร และ Animator จากการทำงานจริง สถานะ “ผ่าน” หมายถึงผ่านกรณีควบคุมที่บันทึกไว้ ไม่ใช่การรับรองว่าการเล่นครบทุกเส้นทางหรือการทดสอบโดยผู้ใช้งานผ่านทั้งหมด'
for p in d.paragraphs:
 if p._p not in b[s:z]:continue
 if p.text.startswith('ตารางต่อไปแยกการตรวจโค้ด'):p.text=scope
 elif p.text.startswith('เพื่อยืนยันผลขณะทำงานจริง กำหนดกรณี'):p.text='ได้ทดสอบ Animator ใน Play Mode ตามกรณีต่อไปนี้ โดยปรับพารามิเตอร์ของ Controller เดิมและตรวจชื่อสถานะที่ทำงานจริง ผลทั้ง 5 กรณีผ่านตามตารางที่ 4-8'
 elif p.text.startswith('ตารางที่ 4-9 รายการที่ยังต้อง'):p.text='ตารางที่ 4-9 สรุปผลทดสอบและสถานะการประเมิน'
 elif 'ทั้งนี้ผลการเปลี่ยนภาพระหว่างเล่นยังต้องยืนยัน' in p.text:
  p.text=p.text.split('ทั้งนี้ผลการเปลี่ยนภาพระหว่างเล่นยังต้องยืนยัน')[0]+'ผลทดสอบ Controller ใน Play Mode ผ่าน AN-P01 และ AN-K01–AN-K04 ตามตารางที่ 4-8'
 elif p.text.startswith('ข้อจำกัดของหลักฐานคือการควบคุมสถานะ'):
  p.text='การทดสอบเพิ่มเติมยืนยันผลในกรณีควบคุมรวม 35 กรณีหลัก ขอบเขตนี้ครอบคลุมองค์ประกอบและการเปลี่ยนสถานะที่กำหนด ยังไม่ใช่การเล่นครบทุกเส้นทาง การประเมินโดยผู้ใช้ หรือการรับรองประสิทธิภาพทุกอุปกรณ์ คะแนนความพึงพอใจยังเป็นตัวอย่างสมมติ'
 elif p.text.startswith('มีต้นแบบเกมสามด่าน พร้อมสคริปต์'):
  p.text=p.text.replace('ส่วนความถูกต้องของการเล่นตั้งแต่เริ่มจนชนะทุกด่านและความพึงพอใจยังต้องทดสอบเพิ่มเติม','ผลทดสอบองค์ประกอบ 35 กรณีผ่านตามขอบเขตในหัวข้อ 4.10 ส่วนการเล่นครบทุกเส้นทางและความพึงพอใจของผู้ใช้ยังต้องประเมินเพิ่มเติม')
 for run in p.runs:
  if run.font.name is None:run.font.name='TH SarabunPSK';run.font.size=Pt(16)
summary=next(t for t in d.tables if 'กลุ่มรายการ'==t.rows[0].cells[0].text)
for c,text in zip(summary.rows[0].cells,['กลุ่มรายการ','ผลที่ตรวจได้','สถานะ']):c.text=text;style_cell(c)
newrows=[('การควบคุมและ Player FSM','TC-01–TC-05 และ TC-FSM-P01–P07 รวม 12 กรณี','ผ่านกรณีควบคุม 12/12'),('Enemy FSM','TC-FSM-E01–E10 รวม 10 กรณี','ผ่านกรณีควบคุม 10/10'),('ระบบเกมและด่าน','TC-06–TC-13 รวม 8 กรณี','ผ่านกรณีควบคุม 8/8'),('แอนิเมชัน','AN-P01 และ AN-K01–AN-K04 รวม 5 กรณี','ผ่านกรณีควบคุม 5/5'),('ผู้ทดลองเล่น','ยังไม่ได้เก็บแบบสอบถามหรือข้อเสนอแนะจริง','ยังไม่ประเมิน; คะแนนเป็นตัวอย่างสมมติ')]
for row,values in zip(summary.rows[1:],newrows):
 for c,text in zip(row.cells,values):c.text=text;style_cell(c)
out=r/'project_บท4_ผลทดสอบจริง_PlayMode.docx';d.save(out)
c=Document(out);bb=list(c.element.body);ss=next(i for i,e in enumerate(bb) if txt(e).startswith('บทที่ 4 ผล'));zz=next(i for i,e in enumerate(bb) if i>ss and txt(e)=='บทที่ 5')
assert before==[etree.tostring(e) for e in bb[:ss]]
assert after==[etree.tostring(e) for e in bb[zz:]]
assert len(descriptions)==35
assert all(evidence[k]['status']=='PASS' for k in descriptions)
print('Updated chapter 4 with 35 passing controlled Play Mode tests; other chapters preserved.')
