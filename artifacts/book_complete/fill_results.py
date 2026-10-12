from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
r=Path(__file__).resolve().parent
d=Document(r/'project_บท3เติมภาพ_บท4ล่าสุด.docx')
tables=list(d.tables)
def fmt(t):
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    for run in p.runs:
     run.font.name='TH SarabunPSK';run.font.size=Pt(14)
     rf=run._element.get_or_add_rPr().get_or_add_rFonts()
     for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'TH SarabunPSK')
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();tag=OxmlElement('w:cantSplit');pr.append(tag)
 hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
for a,b in [(4,8),(5,9),(6,10),(7,11)]:
 tables[a].rows[0].cells[-1].text='ผลที่บันทึกและสถานะ'
 for left,right in zip(tables[a].rows[1:],tables[b].rows[1:]):
  left.cells[-1].text=right.cells[3].text+'\nสถานะ: '+right.cells[4].text
 fmt(tables[a])
 t=tables[a];t.autofit=False
 for tag in ['tblpPr','tblInd']:
  for el in t._tbl.tblPr.findall(qn('w:'+tag)):t._tbl.tblPr.remove(el)
 t._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'),'8352')
 for col,w in zip(t.columns,[.65,1.05,1.1,1.35,1.65]):col.width=Inches(w)
 for row in t.rows:
  for c,w in zip(row.cells,[.65,1.05,1.1,1.35,1.65]):
   c.width=Inches(w)
   for p in c.paragraphs:
    p.paragraph_format.keep_with_next=False
    p.paragraph_format.page_break_before=False
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=1
active=False
for p in list(d.paragraphs):
 if p.text.startswith('3.6 การทดสอบ'):active=True
 if p.text.startswith('3.7 การบำรุง'):active=False
 if active and not p.text.strip() and not p._p.xpath('.//w:drawing'):
  p._p.getparent().remove(p._p)
for p in d.paragraphs:
 if 'ตารางในหัวข้อนี้จึงเว้นผลจริง' in p.text:
  p.text='ตารางการทดสอบบันทึกผลจากการตรวจโค้ดและการสาธิตใน Play Mode ที่มีหลักฐานอยู่แล้ว โดยใช้ข้อมูลเดียวกับบท 4 สถานะสาธิตบางส่วนหมายถึงยังไม่ยืนยันครบทุกเงื่อนไข ส่วนยังไม่ทดสอบหมายถึงตรวจพบโครงสร้างในโค้ดแต่ยังไม่มีผลทดลองครบกรณี จึงไม่ระบุว่าผ่านจนกว่าจะทดสอบครบ'
 if 'ตารางนี้รายงานสถานะของข้อมูลจริง' in p.text:
  p.text=p.text.replace('ตารางนี้รายงานสถานะของข้อมูลจริง ส่วนตัวอย่างคะแนนสมมติแสดงในหัวข้อ 4.11.2','ตารางนี้แสดงค่าเฉลี่ยตัวอย่างจากข้อมูลสมมติ 5 แถวควบคู่กับสถานะของข้อมูลจริง ตัวเลขใช้สาธิตการคำนวณเท่านั้น รายละเอียดแสดงในหัวข้อ 4.11.2')
means=['4.40','4.20','4.00','3.60','4.40','4.40'];levels=['ดีมาก','ดี','ดี','ดี','ดีมาก','ดีมาก']
t=tables[12];t.rows[0].cells[1].text='ค่าเฉลี่ยสมมติ';t.rows[0].cells[2].text='ระดับสมมติและสถานะจริง'
for i,row in enumerate(t.rows[1:]):
 row.cells[1].text=means[i]+' (ตัวอย่าง)'
 row.cells[2].text=levels[i]+' (สมมติ)\nผลจริง: ยังไม่เก็บแบบสอบถาม'
fmt(t)
def insert_table(anchor,caption,headers,rows,widths):
 p=anchor.insert_paragraph_before(caption);p.paragraph_format.keep_with_next=True
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,s in zip(t.rows[0].cells,headers):c.text=s
 for values in rows:
  for c,s in zip(t.add_row().cells,values):c.text=s
 t.autofit=False
 for row in t.rows:
  for c,w in zip(row.cells,widths):c.width=Inches(w)
 anchor._p.addprevious(t._tbl);fmt(t)
 for run in p.runs:run.font.name='TH SarabunPSK';run.font.size=Pt(16)
anchor=next(p for p in d.paragraphs if p.text.startswith('4.10 ผล'))
insert_table(anchor,'ตารางที่ 4-8 ผลตรวจโครงสร้างและสถานะทดสอบแอนิเมชัน', ['รหัส','ผลตรวจโครงสร้าง','สถานะทดสอบจริง'],[
 ['AN-P01','พบ Idle, Walk และ Bool isWalking เชื่อมกับสคริปต์','ตรวจโครงสร้างแล้ว; ยังไม่ทดสอบการเริ่มและหยุดครบกรณี'],
 ['AN-K01','พบ Walk, Run และ Bool isWalking กับ isRunning','ตรวจโครงสร้างแล้ว; ยังไม่ยืนยันภาพเปลี่ยนตาม Patrol และ Chase'],
 ['AN-K02','พบ Trigger attack1 และ Transition ของ Attack1','ตรวจโครงสร้างแล้ว; ยังไม่ยืนยันคลิปโจมตีและการกลับครบลำดับ'],
 ['AN-K03','พบ Trigger scream และ Any State ไป Scream','ตรวจโครงสร้างแล้ว; ยังไม่ยืนยันกรีดร้องจากทุกสถานะ'],
 ['AN-K04','พบ Transition ออกจาก Run ตาม Bool','ตรวจโครงสร้างแล้ว; ยังไม่ยืนยันลำดับเมื่อเงื่อนไขผ่านพร้อมกัน']], [0.7,2.45,2.45])
anchor=next(p for p in d.paragraphs if p.text.startswith('4.11 ผล'))
insert_table(anchor,'ตารางที่ 4-9 รายการที่ยังต้องเก็บผลเพิ่มเติม',['กลุ่มรายการ','สิ่งที่ยังต้องประเมิน','สถานะปัจจุบัน'],[
 ['การควบคุมและ Player FSM','ทดสอบทุกทิศ การกระโดดซ้ำ บันได การคืนสถานะ และแบตเตอรี่','มีหลักฐานสาธิตบางส่วน; ยังไม่สรุปผ่านครบกรณี'],
 ['Enemy FSM','ทดสอบ E02–E03 และ E05–E10 รวมการตรวจพบตามธรรมชาติใน E04','ตรวจโค้ดแล้ว; ยังขาดผลเล่นทดสอบครบลำดับ'],
 ['ระบบเกมและด่าน','ทดสอบการยกเลิกปลดล็อก การสุ่มหลายรอบ เสียงกับดักและหัวใจ','มีภาพและโค้ดบางส่วน; ยังขาดผลครบเงื่อนไข'],
 ['แอนิเมชัน','ตรวจ Play Mode ตาม AN-P01 และ AN-K01–AN-K04','ตรวจ Controller และสคริปต์แล้ว; ยังขาดผลการเปลี่ยนภาพครบลำดับ'],
 ['ผู้ทดลองเล่น','เก็บจำนวนผู้ตอบ คะแนน 6 ด้าน และข้อเสนอแนะ','ยังไม่มีผู้ตอบจริง; คะแนนที่ใส่เป็นตัวอย่างสมมติเท่านั้น']], [1.05,2.55,2.0])
d.save(r/'project_บท3และ4_กรอกตารางผล.docx')
assert not any('รอบันทึกผล' in c.text for t in d.tables for row in t.rows for c in row.cells)
print('Filled chapter 3 evidence, completed evaluation example cells, added 2 result tables.')
