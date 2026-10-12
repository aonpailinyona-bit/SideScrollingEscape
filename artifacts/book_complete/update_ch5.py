from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree
r=Path(__file__).resolve().parent
d=Document(r/'project_บท3และ4_ผลทดสอบล่าสุด.docx');p=list(d.paragraphs)
b=list(d.element.body);txt=lambda e:''.join(e.xpath('.//w:t/text()'))
s=next(i for i,e in enumerate(b) if txt(e)=='บทที่ 5');before=[etree.tostring(e) for e in b[:s]]
def insert(anchor,text,template):
 n=anchor.insert_paragraph_before(text)
 if n._p.pPr is not None:n._p.remove(n._p.pPr)
 if template._p.pPr is not None:n._p.insert(0,deepcopy(template._p.pPr))
 for run in n.runs:
  if template.runs and template.runs[0]._r.rPr is not None:run._r.insert(0,deepcopy(template.runs[0]._r.rPr))
 return n
insert(p[551],'ผลทดสอบทางเทคนิคใน Unity 6000.4.9f1 เมื่อวันที่ 4 ตุลาคม 2569 ผ่านครบ 35 กรณีหลัก แบ่งเป็นระบบควบคุมพื้นฐาน 5 กรณี Player FSM 7 กรณี Enemy FSM 10 กรณี ระบบเกมและด่าน 8 กรณี และแอนิเมชัน 5 กรณี รายละเอียดผลระบบ 30 กรณีอยู่ในหัวข้อ 3.6 และ 4.10 ส่วนผลแอนิเมชันอยู่ในตารางที่ 4-8 การทดสอบใช้ Play Mode แบบควบคุมอินพุต ตำแหน่ง เวลา และเงื่อนไขขององค์ประกอบ จึงยืนยันเฉพาะกรณีที่กำหนดไว้ ยังไม่ใช่ผลการเล่นครบทุกเส้นทางหรือการประเมินโดยผู้ใช้งาน',p[549])
insert(p[572],'5.1.7 ผลการพัฒนาระบบแอนิเมชัน',p[569])
insert(p[572],'ผู้เล่นใช้คลิป Idle และ Walk โดยพารามิเตอร์ Bool isWalking ควบคุมการเปลี่ยนภาพตามการเคลื่อนที่ ส่วน Killer ใช้ Idle, Walk, Run, Attack1 และ Scream โดยใช้ Bool isWalking และ isRunning รักษาสถานะการเดินและวิ่ง และใช้ Trigger attack1 กับ scream เริ่มการโจมตีและกรีดร้องเฉพาะครั้ง FSM ตัดสินใจพฤติกรรม ขณะที่ Animator แสดงภาพของพฤติกรรมนั้น',p[570])
insert(p[572],'ผลทดสอบ Controller เดิมใน Play Mode ผ่าน AN-P01 และ AN-K01–AN-K04 รวม 5 กรณี ได้แก่ Idle → Walk → Idle ของผู้เล่น การเดินและวิ่งของ Killer การเล่น Attack1 แล้วกลับ Idle การเข้าสู่ Scream จาก Idle, Walk, Run และ Attack1 แล้วกลับ Idle และการออกจาก Run ไป Walk เมื่อ isRunning เปลี่ยนเป็น false ทั้งนี้ Float และ Int ยังไม่ได้ใช้ใน Controller ปัจจุบัน จึงเป็นแนวทางขยายระบบตามหัวข้อ 3.5.11 และไม่มีผลทดสอบของพารามิเตอร์สองชนิดนี้',p[570])
p[572].text='5.1.8 อภิปรายผลการพัฒนา'
insert(p[577],'ผลที่ผ่านทั้ง 35 กรณีสนับสนุนว่าการเปลี่ยนสถานะและการเชื่อม Animator ทำงานได้ตามเงื่อนไขควบคุมที่กำหนด อย่างไรก็ตาม การทดสอบใช้การเรียกองค์ประกอบและจัดสภาพแวดล้อมผ่านสคริปต์ชั่วคราว จึงยังต้องตรวจการทำงานร่วมกันระหว่างเล่นครบเส้นทาง ความเสถียรเมื่อเล่นต่อเนื่อง และการตอบสนองบนอุปกรณ์ต่างกันก่อนสรุปความพร้อมสำหรับเผยแพร่',p[573])
insert(p[577],'ยังไม่ได้เก็บแบบสอบถามจากผู้ทดลองเล่น จึงไม่สามารถสรุประดับความพึงพอใจ ความสมดุลของความยาก หรือความน่ากลัวที่ผู้เล่นรับรู้ได้ คะแนน ค่าเฉลี่ย และส่วนเบี่ยงเบนมาตรฐานตัวอย่างในบท 4 เป็นข้อมูลสมมติสำหรับแสดงวิธีคำนวณเท่านั้น ไม่ใช่หลักฐานผลประเมินจริงและไม่ใช้สนับสนุนข้อสรุปด้านประสบการณ์ผู้เล่น',p[573])
p[605].text='ก่อนนำเกมไปใช้งานหรือเผยแพร่ ควรต่อยอดจากผลทดสอบองค์ประกอบ 35 กรณีที่ผ่านแล้ว ด้วยการเล่นครบเส้นทางตั้งแต่เริ่มจนจบทั้งสามด่าน ทดสอบการทำงานร่วมกันของผู้เล่น ศัตรู ไอเท็ม ประตู UI และเสียง รวมถึงกรณีขอบเขต เช่น การเปลี่ยนสถานะพร้อมกัน การยกเลิกปลดล็อกในหลายจังหวะ และการเล่นซ้ำต่อเนื่อง โดยบันทึกผลและข้อผิดพลาดจริงแยกจากผลทดสอบแบบควบคุม'
p[617].text='ควรเก็บข้อมูลจากผู้ทดลองเล่นจริง โดยคัดเลือกผู้เล่นใหม่และผู้ที่คุ้นเคยกับเกมแนวสยองขวัญ กำหนดขั้นตอนเล่นให้เหมือนกัน ระบุจำนวนผู้ตอบ ช่วงเวลาทดสอบ คะแนนทั้ง 6 ด้าน และข้อเสนอแนะ จากนั้นคำนวณค่าเฉลี่ยและส่วนเบี่ยงเบนมาตรฐานจากข้อมูลจริง เพื่อใช้ปรับระดับความยาก การควบคุม และการนำเสนอข้อมูลบนหน้าจอ'
for item in [p[572],p[605],p[617]]:
 for run in item.runs:
  run.font.name='TH SarabunPSK';run.font.size=Pt(16);run.bold=item is p[572]
  rf=run._element.get_or_add_rPr().get_or_add_rFonts()
  for key in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+key),'TH SarabunPSK')
active=False
for paragraph in list(d.paragraphs):
 if paragraph.text=='บทที่ 5':active=True
 if not active:continue
 if paragraph.text.startswith('5.'):
  paragraph.paragraph_format.keep_with_next=True
 if not paragraph.text.strip() and not paragraph._p.xpath('.//w:drawing') and not paragraph._p.xpath('.//w:sectPr'):
  paragraph._p.getparent().remove(paragraph._p)
out=r/'project_บท3ถึง5_ปรับผลทดสอบล่าสุด.docx';d.save(out)
c=Document(out);bb=list(c.element.body);ss=next(i for i,e in enumerate(bb) if txt(e)=='บทที่ 5');assert before==[etree.tostring(e) for e in bb[:ss]]
print('Chapter 5 updated; all earlier chapters preserved exactly.')
