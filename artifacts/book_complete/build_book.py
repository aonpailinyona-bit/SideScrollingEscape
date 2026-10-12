from pathlib import Path
import re, math, statistics
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parent
D=Document(next((R.parent/'animation_ch3').glob('*3และ4*.docx')))
P=list(D.paragraphs)
OUT=R/'project_ฉบับสมบูรณ์พร้อมภาพ.docx'
def text(i,s):
    P[i].text=s
def remove(p):
    if p._p.getparent() is not None:p._p.getparent().remove(p._p)
def font(run,size=16,bold=False):
    run.font.name='TH SarabunPSK';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for k in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+k),'TH SarabunPSK')
    cs=OxmlElement('w:szCs');cs.set(qn('w:val'),str(size*2));run._element.get_or_add_rPr().append(cs)
def add(s='',style=None):
    p=D.add_paragraph(s,style);return p
def before(anchor,s='',style=None):return anchor.insert_paragraph_before(s,style)
def field(p,instruction):
    r=p.add_run(); b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
    r=p.add_run();t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+instruction+' ';r._r.append(t)
    r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'end');r._r.append(b)

# Correct scope and conclusions against the inspected implementation.
text(47,'โครงงานนี้มีวัตถุประสงค์เพื่อพัฒนาเกมสองมิติแนวเอาชีวิตรอดและลอบเร้นด้วย Unity และภาษา C# โดยใช้เครื่องสถานะจำกัด (Finite State Machine: FSM) ควบคุมตัวละครผู้เล่นและศัตรู Killer ระบบผู้เล่นประกอบด้วย Normal, Hidden, Frozen, Hurt, UsingItem, Transitioning และ Dead ส่วนศัตรูใช้สถานะสำหรับลาดตระเวน รับรู้เสียง ตรวจสอบ ค้นหา ไล่ล่า โจมตี และกรีดร้อง การตรวจจับใช้ Raycasting ร่วมกับเหตุการณ์เสียง ผู้เล่นต้องเก็บกุญแจให้ครบและปลดล็อกประตูเพื่อผ่านด่านทั้งสาม ระบบ Animator เชื่อมพฤติกรรมกับภาพเคลื่อนไหวผ่านพารามิเตอร์ Bool และ Trigger ผลจากการตรวจโค้ดและสาธิตใน Play Mode พบโครงสร้างระบบและหน้าจอแสดงผลตามสถานะที่กำหนด อย่างไรก็ตาม การสาธิตมีการควบคุมตำแหน่งและเรียกสถานะผ่านสคริปต์ จึงยังไม่ยืนยันการเล่นครบวงจรหรือความถูกต้องทุกกรณี ยังไม่ได้เก็บแบบสอบถามจากผู้ทดลองเล่นและไม่สามารถสรุประดับความพึงพอใจจริงได้ รายงานแสดงแผนทดสอบ ข้อจำกัด และตัวอย่างข้อมูลสมมติสำหรับอธิบายวิธีวิเคราะห์คะแนนโดยแยกจากผลจริงอย่างชัดเจน')
text(65,P[65].text.replace('เกมสยองขวัญ 2 มิติเชิงจริยธรรม','เกมสยองขวัญ 2 มิติ'))
text(66,'ขอกราบขอบพระคุณบิดา มารดา และครอบครัว ที่สนับสนุนการศึกษาและการจัดทำโครงงาน รวมถึงผู้ที่ให้คำแนะนำและความช่วยเหลือในการพัฒนาระบบ ทั้งนี้การเก็บข้อมูลประเมินความพึงพอใจจากผู้ทดลองเล่นยังเป็นขั้นตอนที่จะดำเนินการต่อไป')
text(84,'เกมแนวลอบเร้นต้องควบคุมการเคลื่อนที่ การซ่อนตัว การใช้ไอเท็ม และการตอบสนองต่ออันตรายให้สัมพันธ์กัน หากตรวจเงื่อนไขทั้งหมดโดยไม่มีโครงสร้างสถานะ อาจเกิดคำสั่งที่ขัดแย้ง เช่น การเคลื่อนที่ขณะซ่อนหรือเสียชีวิต โครงงานจึงใช้ FSM แยกพฤติกรรมและเงื่อนไขการรับคำสั่งของผู้เล่นและศัตรูให้ตรวจสอบได้ชัดเจน')
text(89,'1.2.2 เพื่อประยุกต์ใช้ Finite State Machine ในการจัดการสถานะการควบคุมตัวละครผู้เล่น')
text(90,'1.2.3 เพื่อพัฒนาปัญญาประดิษฐ์ของศัตรูที่ตอบสนองต่อการมองเห็น เสียง และสถานะการซ่อนของผู้เล่น')
text(97,'ใช้ FSM ควบคุมสถานะ Normal, Hidden, Frozen, Hurt, UsingItem, Transitioning และ Dead โดยแยกจากสถานะภาพเคลื่อนไหว Idle และ Walk ใน Animator')
text(98,'1.3.1.2 การควบคุมการชนและการซ่อนตัว')
text(99,'ใช้ Rigidbody2D และ Collider2D ควบคุมการเคลื่อนที่ การตรวจพื้น และการชน ส่วน HidingSystem ควบคุมการเข้าและออกจากจุดซ่อน รุ่นที่ตรวจสอบยังไม่มีการนั่งย่อและหมอบคลานเป็นสถานะผู้เล่น')
text(108,'1.3.2.3 การเคลื่อนที่ตามเป้าหมาย')
text(109,'ศัตรูเคลื่อนที่แนวนอนด้วย Rigidbody2D ไปยังจุดลาดตระเวน ตำแหน่งเสียง หรือตำแหน่งผู้เล่น พร้อมตรวจพื้นและสิ่งกีดขวาง รุ่นปัจจุบันยังไม่ได้ใช้ NavMesh หรือ A* สำหรับคำนวณเส้นทาง')
text(113,'ผู้เล่นต้องเก็บกุญแจ 5 ดอกตาม LevelConfig ของแต่ละด่าน แล้วปลดล็อกประตูทางออก')
text(115,'ใช้ GameManager และ SceneManager จัดการเมนู ด่าน การแพ้และการชนะ โดยการบันทึกความคืบหน้าข้ามการเปิดเกมเป็นข้อเสนอสำหรับพัฒนาต่อ')
text(119,'1.3.5.1 เกมสองมิติแนวลอบเร้น (Stealth Game) โดยผู้เล่นหลบศัตรูและวางแผนการเคลื่อนที่')
text(120,'1.3.5.2 มีฉากการเล่น 3 ด่าน พร้อมสิ่งกีดขวาง จุดซ่อน ไอเท็ม และประตูทางออก')
text(125,'1.3.5.6 ผู้เล่นแพ้เมื่อพลังชีวิตหมด และระบบแสดงหน้าจอ Game Over')
text(130,'1.3.6.4 ใช้ Finite State Machine ควบคุมสถานะของตัวละครผู้เล่นและศัตรู')
text(132,'1.3.6.6 ใช้การเคลื่อนที่ด้วย Rigidbody2D และการตรวจพื้นกับสิ่งกีดขวาง โดยศึกษาการค้นหาเส้นทางเป็นแนวทางสำหรับพัฒนาต่อ')
text(138,'1.5 แผนการดำเนินโครงงาน')
text(140,'แผนงานแบ่งเป็นการศึกษาข้อมูล วิเคราะห์และออกแบบ พัฒนาโมดูล ทดสอบระบบ และจัดทำรายงาน ตารางกำหนดเวลาในต้นฉบับใช้เป็นแผนการดำเนินงาน ไม่ใช้ยืนยันวันที่เสร็จจริงของแต่ละกิจกรรม')
text(141,'สถานะปัจจุบันคือมีต้นแบบเกมและรายงานประกอบ การตรวจกรณีทดสอบครบวงจรและการเก็บข้อมูลผู้ทดลองเล่นยังต้องดำเนินการต่อ')
text(142,'1.6 ประโยชน์ที่คาดว่าจะได้รับ')
text(147,'1.6.5 ได้แนวทางเก็บข้อมูลความพึงพอใจเพื่อใช้ปรับการควบคุม ความยาก และการสื่อสารข้อมูลของเกม')
text(162,'การพัฒนาเกมเป็นกระบวนการสร้างซอฟต์แวร์เชิงโต้ตอบ ประกอบด้วยการวิเคราะห์ความต้องการ ออกแบบกลไก พัฒนาโปรแกรมและสื่อ ทดสอบ และประเมินผล องค์ประกอบหลัก ได้แก่ การควบคุม กราฟิก แอนิเมชัน เสียง ปัญญาประดิษฐ์ และส่วนติดต่อผู้ใช้ การแยกองค์ประกอบเป็นโมดูลช่วยให้ตรวจสอบปัญหาและปรับระบบได้สะดวก โดยควรใช้ผลการทดสอบจริงยืนยันว่าระบบที่พัฒนาตอบสนองต่อเป้าหมายของเกม')
text(164,'เกมสองมิติใช้พิกัดแกน X และ Y ในการกำหนดตำแหน่ง โดย Unity รองรับ Sprite, Tilemap, Physics 2D และ Animator การออกแบบต้องพิจารณาระดับพื้น สิ่งกีดขวาง การชน และขอบเขตกล้อง เพื่อให้ภาพกับพื้นที่ที่ผู้เล่นเคลื่อนที่ได้สอดคล้องกัน ประสิทธิภาพขึ้นอยู่กับจำนวนวัตถุ เอฟเฟกต์ การตรวจฟิสิกส์ และวิธีจัดการทรัพยากร จึงต้องวัดบนอุปกรณ์เป้าหมาย')
text(171,'เกมแนวซ่อนตัวและสยองขวัญให้ผู้เล่นหลบผู้ไล่ล่า รวบรวมสิ่งของ และหาทางออก โดยใช้การมองเห็น เสียง และทรัพยากรจำกัดสร้างความเสี่ยง ในโครงงานนี้ผู้เล่นเลือกเดินหรือวิ่ง ใช้ไฟฉาย ซ่อนตัว และใช้ก้อนหินสร้างเสียงล่อ การปรับพฤติกรรมศัตรูและการจัดวางฉากมีผลต่อความยาก จึงควรประเมินจากทั้งการทดสอบทางเทคนิคและผู้ทดลองเล่น')
text(174,'องค์ประกอบของ FSM ได้แก่ สถานะเริ่มต้น สถานะที่เป็นไปได้ เหตุการณ์ เงื่อนไข และการเปลี่ยนสถานะ โครงงานใช้ FSM ในสคริปต์เพื่อกำหนดสิทธิ์การควบคุมผู้เล่นและพฤติกรรมศัตรู ส่วน Animator ใช้เครื่องสถานะอีกชุดหนึ่งสำหรับเลือกภาพเคลื่อนไหว ทั้งสองส่วนเชื่อมกันผ่านพารามิเตอร์ จึงไม่จำเป็นต้องมีชื่อหรือจำนวนสถานะเท่ากัน เช่น Normal รองรับทั้งการยืนและเดิน ส่วน Chase ของ Killer ส่งค่าให้เล่น Run')
text(179,'2.6.3 การเคลื่อนที่และการนำทาง ระบบนำทางอาจใช้กราฟ กริด หรือข้อมูลพื้นผิวร่วมกับอัลกอริทึมค้นหาเส้นทาง ในต้นแบบนี้ KillerAI กำหนดความเร็วแนวนอนตามตำแหน่งเป้าหมายและตรวจสิ่งกีดขวาง การนำ A* มาใช้เป็นแนวทางพัฒนาต่อสำหรับฉากที่ซับซ้อนกว่าเดิม')
text(196,'การค้นหาเส้นทางเป็นการเลือกเส้นทางจากจุดเริ่มต้นไปยังเป้าหมายบนข้อมูลกราฟหรือกริด A* จัดลำดับโหนดด้วย f(n) = g(n) + h(n) โดย g(n) คือค่าใช้จ่ายที่สะสมและ h(n) คือค่าประมาณที่เหลือ คุณสมบัติการหาเส้นทางเหมาะสมขึ้นอยู่กับเงื่อนไขของ heuristic และวิธีใช้งาน หัวข้อนี้เป็นความรู้สำหรับพัฒนาต่อ รุ่นปัจจุบันใช้การเคลื่อนที่แนวนอนและตรวจพื้น ไม่ได้เรียก A* หรือ NavMesh จึงยังไม่มีผลเปรียบเทียบประสิทธิภาพของอัลกอริทึมในโครงงาน')
text(197,'2.9 การออกแบบ UX และ UI ให้เหมาะสม')
text(198,'UI ต้องสื่อสารทรัพยากรและอันตรายโดยไม่บดบังพื้นที่เล่น โครงงานใช้ HUD แสดง Health, Stamina, Battery, Keys และจำนวนก้อนหิน พร้อมข้อความ HIDDEN และ DANGER และ Prompt สำหรับวัตถุที่โต้ตอบได้ สีและตำแหน่งข้อความควรอ่านง่าย ส่วนเสียงและขอบจอช่วยแจ้งอันตรายโดยไม่ต้องมองเห็นศัตรูตลอดเวลา ความชัดเจนของการสื่อสารต้องตรวจด้วยผู้ทดลองเล่นเพิ่มเติม')
text(199,'2.10 งานวิจัยและเอกสารที่เกี่ยวข้อง')
related=[
('Finite State Machine in Game Development','Jagdale (2021) อธิบายการจัดพฤติกรรม NPC ด้วยสถานะและเงื่อนไขการเปลี่ยนสถานะ รวมถึงแนวคิด FSM แบบลำดับชั้น แนวคิดนี้ใช้เป็นพื้นฐานในการแยกการลาดตระเวน การตรวจสอบ และการโจมตีของ Killer โดยรายงานนี้ไม่ใช้บทความดังกล่าวเป็นหลักฐานว่าต้นแบบมีเฟรมเรตหรือความพึงพอใจสูงขึ้น'),
('Implementation of Inventory Checking System Using Finite State Machine in Role Playing Game','Kusuma และ Handayani (2024) ใช้ FSM ร่วมกับระบบตรวจสอบไอเท็มและภารกิจใน RPG Maker MV โดยพัฒนาตามวงจร GDLC แนวคิดที่นำมาประยุกต์คือการตรวจเงื่อนไขไอเท็มก่อนเปลี่ยนสถานะภารกิจ ซึ่งสัมพันธ์กับการตรวจจำนวนกุญแจก่อนเปลี่ยน ExitDoor จาก Locked เป็น Ready แม้เกมเอนจินและประเภทเกมต่างกัน'),
('An Analysis of Finite State Machine Based Enemy Artificial Intelligence in Kirby Nightmare in Dream Land','Thyrun (2026) วิเคราะห์แบบจำลองพฤติกรรมศัตรูและผลของเวลาคูลดาวน์ต่อความยากด้วยการสังเกตและการจำลอง บทความเป็น preprint ที่ยังไม่ผ่านการประเมินโดยผู้ทรงคุณวุฒิ ใช้ประกอบแนวคิดการแบ่งสถานะโจมตีและพักหลังโจมตี ไม่ใช้เป็นผลทดสอบของ Killer ในโครงงานนี้'),
('Unity Animator Controller','เอกสาร Unity Technologies อธิบายการใช้ Animator Controller เป็นแผนผังสถานะและ Transition เพื่อเลือกคลิปตามพารามิเตอร์ ใช้เป็นแนวทางสร้าง Idle และ Walk ของผู้เล่น และ Idle, Walk, Run, Attack1 และ Scream ของ Killer โดยแยกการแสดงภาพออกจากตรรกะ FSM ในสคริปต์'),
('Unity Animation Parameters','เอกสาร Unity Technologies แบ่งพารามิเตอร์เป็น Float, Integer, Bool และ Trigger การเลือกใช้ขึ้นอยู่กับลักษณะข้อมูล โดย Bool เหมาะกับเงื่อนไขที่คงค่า เช่น กำลังเดิน ส่วน Trigger เหมาะกับเหตุการณ์เริ่มการกระทำ เช่น โจมตี Float และ Int เป็นทางเลือกสำหรับค่าต่อเนื่องและรหัสสถานะ ซึ่งต้นแบบยังไม่ได้เพิ่มใน Animator'),
('Unity Physics2D Raycast','เอกสาร Physics2D.Raycast อธิบายการส่งรังสีจากตำแหน่งต้นทางไปตามทิศทางและระยะที่กำหนด พร้อมกรองเลเยอร์ ใช้ประกอบการตรวจสิ่งกีดขวางในแนวสายตาของ Killer การตรวจระยะและมุมเกิดก่อนส่ง Raycast ตามโค้ดของโครงงาน ไม่ใช้ Raycast เป็นหลักฐานว่าไม่มีข้อผิดพลาดในการตรวจจับทุกตำแหน่ง'),
('Improving Mobile Game Performance with Basic Optimization Techniques in Unity','Koulaxidis และ Xinogalos (2022) ศึกษาการปรับประสิทธิภาพเกม Unity บนอุปกรณ์มือถือโดยเปรียบเทียบรุ่นที่ใช้เทคนิคปรับทรัพยากรและการจัดการวัตถุ ข้อที่นำมาใช้คือควรวัดประสิทธิภาพบนอุปกรณ์จริงก่อนสรุปผล โครงงานนี้ยังไม่ได้เก็บเฟรมเรตอย่างเป็นระบบ และไม่ย้ายผลตัวเลขของเกมในบทความมาใช้เป็นผลของต้นแบบ')]
for j,(title,body) in enumerate(related):text(200+j*2,f'2.10.{j+1} {title}');text(201+j*2,body)
for i in [232,317,338,543]:
    s=P[i].text.replace('ป่ามืด','สวนสนุกและพื้นที่ป่ามืด').replace('เป็นป่า โดย','เป็นสวนสนุกที่มีพื้นที่ป่า โดย');text(i,s)
text(237,'ศึกษา FSM เพื่อแบ่งสถานะและเงื่อนไขของผู้เล่น ศัตรู ประตูทางออก และระดับความตึงเครียด ส่วน InventoryManager จัดการข้อมูลไอเท็มและเหตุการณ์แจ้ง UI ระบบเสียงและคำเตือนตอบสนองต่อข้อมูลสถานะและระยะห่าง')
text(287,'เมื่อเข้าใกล้วัตถุที่โต้ตอบได้ ระบบแสดง Prompt และใช้ปุ่ม E เพื่อสั่งงาน เช่น ประตูทางออก ส่วน ItemPickup เก็บกุญแจเมื่อผู้เล่นสัมผัส Trigger ของวัตถุ')
text(299,'ก้อนหินใช้สร้างเหตุการณ์เสียงประเภท Rock จากตำแหน่งอื่นเพื่อเบี่ยงเบนศัตรู ภาพสาธิตแสดงการส่งเหตุการณ์ Rock ผ่าน HearNoise แล้ว Killer เข้าสู่ Listening ยังไม่ได้ยืนยันลำดับการขว้าง การตก และการเดินไปตรวจสอบครบวงจร')
text(301,'ภาพที่ 3-9 สาธิต Killer รับเหตุการณ์เสียงประเภท Rock')
text(289,'ภาพที่ 3-5 สาธิตตำแหน่งกุญแจและระบบเก็บไอเท็ม')
text(308,'เมื่อกุญแจครบ ผู้เล่นกด E ค้างที่ประตูจนความคืบหน้าครบ การปลดล็อกยกเลิกเมื่อปล่อยปุ่ม ออกจากระยะ ถูกโจมตี หรือมีศัตรูอยู่ใกล้ ภาพนี้เรียกสถานะ Unlocking และกำหนดความคืบหน้าเพื่อแสดง Prompt ของระบบ')
text(310,'ภาพที่ 3-12 สาธิตข้อความและความคืบหน้าการปลดล็อกประตู')
text(319,'ภาพที่ 3-15 สภาพแวดล้อมสวนสนุกและพื้นที่ป่าใน Level 3')
text(322,'ภาพที่ 3-16 สาธิตตำแหน่งกับดักเสียงใน Level 3')
text(328,P[328].text.replace('AudioManager และ TensionDirector','AudioManager, LevelAudioController และ Level3HeartbeatWarning'))
text(334,'TensionDirector เป็นคลาสที่ประกาศในไฟล์ LevelAudioController.cs ใช้สถานะ Calm, Suspicious, Hunted, Escaped และ Cooldown โดยรวมพฤติกรรมศัตรูทุกตัวและส่งเหตุการณ์ให้ระบบเสียง ระหว่าง Hunted ปรับตัวคูณความเร็วไล่ล่า ส่วน UIManager แสดง DANGER จากสถานะ Killer และ Level3HeartbeatWarning ใช้ระยะห่างควบคุมเสียงหัวใจกับขอบจอสีแดง จึงต้องแยก FSM ความตึงเครียดออกจากการแสดงคำเตือนโดยตรง')
text(356,'แผนภาพสรุปเส้นทางหลักของ TensionDirector และการส่งข้อมูลให้ระบบเสียง โดย Hunted เกิดเมื่อมีศัตรูล่า ผู้เล่นหลุดจากการล่าจะเข้าสู่ Escaped แล้ว Cooldown ก่อนกลับ Calm หากมีการล่าใหม่จะเข้าสู่ Hunted จากสถานะใดก็ได้')
text(358,'ภาพที่ 3-21 แผนภาพสถานะความตึงเครียดและการส่งข้อมูลเสียง')
text(398,'กำหนด Test Case ด้วยขั้นตอนและผลที่คาดหวัง ตารางบทนี้เป็นแผนทดสอบ สถานะ “กำหนดกรณีแล้ว” หมายถึงมีขั้นตอนพร้อมใช้ ยังไม่ใช่ผลผ่าน การตรวจโค้ดและสาธิตแบบควบคุมสถานะใน Play Mode แสดงแยกไว้ในบท 4')
text(396,'ตรวจสอบจาก PlayerAnimator.controller, KillerAnimator.controller, PlayerController.cs และ KillerAI.cs ในโครงการ SideScrollingEscape ส่วนชนิดพารามิเตอร์อ้างอิง Unity Manual เรื่อง Animation Parameters ตามบรรณานุกรม ตรวจสอบวันที่ 4 ตุลาคม 2569')
text(439,'บทนี้นำเสนอผลการตรวจโครงสร้างและการสาธิตต้นแบบใน Unity Play Mode เมื่อวันที่ 4 ตุลาคม 2569 ภาพเกมจับจากการทำงานจริง ส่วนภาพเฟรมแอนิเมชันและแผนภาพระบุแหล่งที่มาตามคำบรรยาย การสาธิตใช้สคริปต์ชั่วคราวควบคุมตำแหน่ง เรียกเมธอดและบางสถานะ หยุดการอัปเดตศัตรูเพื่อจัดภาพ และปิดโอกาสสุ่ม Jumpscare ในช่วงเก็บภาพ จึงไม่ใช่การทดลองเล่นครบวงจรและไม่ยืนยันผลผ่านทุก Test Case ยังไม่มีข้อมูลแบบสอบถามจากผู้ทดลองเล่น')
text(448,'บันทึกการสาธิตพบสถานะ Normal, UsingItem, Hidden, Frozen, Hurt และ Dead ตามเมธอดที่เรียก ภาพแสดงผล HIDDEN และ Game Over สอดคล้องกับสถานะที่บันทึก การสาธิตนี้ยืนยันการแสดงผลภายใต้เงื่อนไขที่จัดไว้ ยังไม่ยืนยันว่าทุกการเปลี่ยนสถานะจะถูกต้องเมื่อหลายเหตุการณ์เกิดพร้อมกัน')
text(450,'ภาพที่ 4-3 สาธิตหน้าจอ Normal และ Hidden ของผู้เล่น')
text(454,'ภาพที่ 4-4 Killer ในสถานะ Patrol ที่จัดไว้เพื่อสาธิต')
text(456,'ภาพที่ 4-5 สาธิตการรับเสียง Rock และสถานะ Listening')
text(458,'ภาพที่ 4-6 สาธิตสถานะ Chase และข้อความ DANGER')
text(464,'ภาพที่ 4-7 แผนภาพเงื่อนไขการตรวจแนวสายตาจาก KillerAI')
text(466,'ภาพที่ 4-8 สาธิตพื้นที่กับดักเสียงของ Level 3')
text(487,'ภาพที่ 4-14 สวนสนุกและพื้นที่ป่าใน Level 3')
text(489,'TensionDirector รวมสถานะของศัตรูเป็น Calm, Suspicious, Hunted, Escaped และ Cooldown แล้วส่งเหตุการณ์ให้ LevelAudioController ปรับดนตรี ส่วน UIManager แสดง DANGER และ Level3HeartbeatWarning ควบคุมคำเตือนตามระยะห่าง ภาพเปรียบเทียบแสดงหน้าจอปกติกับหน้าจอที่เรียก Chase เพื่อสาธิต ไม่ได้จับหน้าต่าง Debug ของ TensionDirector จึงยังไม่ใช้ภาพยืนยันลำดับของทั้งห้าสถานะครบถ้วน')
text(490,'สิ่งที่เห็นจากภาพคือข้อความ DANGER และขอบจอสีแดงตามเงื่อนไขที่จัดไว้ ส่วนความชัดเจนของเสียงและความเหมาะสมของระดับความกดดันต้องตรวจด้วยการฟังและทดลองเล่นเพิ่มเติม')
text(492,'ภาพที่ 4-15 เปรียบเทียบหน้าจอปกติกับคำเตือนเมื่อเรียก Chase')
text(529,'ตารางต่อไปแยกการตรวจโค้ดและการสาธิตออกจากการทดสอบครบวงจร “สาธิตบางส่วน” หมายถึงสังเกตผลบางส่วนใน Play Mode ภายใต้การควบคุมของสคริปต์ ส่วน “ยังไม่ทดสอบ” หมายถึงยังไม่มีหลักฐานเพียงพอสำหรับผลที่คาดหวังทั้งแถว จึงยังไม่มีแถวใดสรุปว่าผ่านครบทุกเงื่อนไข')
remove(P[537])
text(539,'กำหนดการประเมิน 6 ด้านด้วยคะแนน 1–5 โดย 5 หมายถึงดีมาก 4 ดี 3 ปานกลาง 2 น้อย และ 1 น้อยมาก ณ วันที่จัดทำรายงานยังไม่ได้เก็บแบบสอบถาม จึงไม่มีจำนวนผู้ตอบจริง ค่าเฉลี่ย ส่วนเบี่ยงเบนมาตรฐาน หรือข้อเสนอแนะจากผู้ใช้ ตารางนี้รายงานสถานะของข้อมูลจริง ส่วนตัวอย่างคะแนนสมมติแสดงในภาคผนวก ข')
text(540,'ตารางที่ 4-5 สถานะข้อมูลการประเมินจากผู้ใช้งานจริง')
text(541,'ยังไม่สามารถจัดอันดับด้านที่ได้รับคะแนนสูงสุดหรือสรุประดับความพึงพอใจจริงได้ หลังเก็บข้อมูลควรคำนวณค่าเฉลี่ยและส่วนเบี่ยงเบนมาตรฐานรายด้าน ระบุจำนวนผู้ตอบและช่วงเวลาทดสอบ แล้วปรับบทอภิปรายให้สอดคล้องกับข้อมูลที่ได้ ตัวเลขในภาคผนวก ข เป็นตัวอย่างสมมติและไม่ใช้แทนผลส่วนนี้')
text(543,'มีต้นแบบเกมสามด่าน พร้อมสคริปต์ควบคุมผู้เล่น Killer ไอเท็ม ประตู และส่วนติดต่อผู้ใช้ การสาธิตแสดงหน้าจอเมนู การเคลื่อนที่ คลังไอเท็ม สถานะซ่อน คำเตือน และหน้าจอผลลัพธ์ภายใต้การจัดเงื่อนไข การตรวจไฟล์ยืนยันโครงสร้าง FSM และการเชื่อม Animator ส่วนความถูกต้องของการเล่นตั้งแต่เริ่มจนชนะทุกด่านและความพึงพอใจยังต้องทดสอบเพิ่มเติม')
text(545,'ข้อจำกัดของหลักฐานคือการควบคุมสถานะระหว่างถ่ายภาพและยังไม่มีผู้ทดลองเล่น รายงานจึงสรุปผลเฉพาะโครงสร้างกับการแสดงผลที่ตรวจพบ ไม่รายงานอัตราผ่าน เวลาเล่น หรือคะแนนความพึงพอใจเป็นผลจริง')
text(606,P[606].text.replace('และ TensionDirector','และระบบเสียงของด่าน'))
text(606,P[606].text+' ระบบความตึงเครียดยังใช้ TensionDirector ที่ประกาศในไฟล์ LevelAudioController.cs เพื่อรวมสถานะศัตรูและแจ้งระบบเสียง')
text(647,'ควรเริ่มเก็บข้อมูลจากผู้ทดลองเล่นที่มีประสบการณ์ต่างกัน ระบุวิธีคัดเลือก จำนวนผู้ตอบ วันที่และอุปกรณ์ทดสอบ พร้อมเก็บข้อมูลดิบเพื่อทวนสอบ เนื่องจากรุ่นรายงานนี้ยังไม่มีข้อมูลจากกลุ่มตัวอย่างจริง')
for i in [609,613,616,620,623,625,628]:
    text(i,P[i].text.replace('พบว่า','มีความเสี่ยงว่า').replace('ในระหว่างการพัฒนา พบ','ประเด็นที่ต้องเฝ้าระวังคือ'))

# Append explicit conclusion against objectives, avoiding unsupported completion claims.
q=before(P[607],'5.1.8 สถานะการบรรลุวัตถุประสงค์', 'Heading 3')
before(P[607],'ด้านการพัฒนาซอฟต์แวร์ มีต้นแบบและโครงสร้าง FSM ของผู้เล่นและศัตรู พร้อม Animator ที่อธิบายในบท 3 และผลตรวจในบท 4 ด้าน Gameplay Loop มีเมธอดและหน้าจอผ่านด่าน แต่ยังต้องเล่นทดสอบต่อเนื่องทุกด่านเพื่อยืนยัน ด้านการประเมินความพึงพอใจยังไม่บรรลุ เนื่องจากยังไม่ได้เก็บแบบสอบถาม ตัวอย่างข้อมูลสมมติในภาคผนวกใช้สอนวิธีคำนวณเท่านั้น')

# Diagrams drawn from source, clearly distinguished from screenshots.
F=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',28)
FS=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',22)
def diagram(name,title,rows):
    w=1500; h=115+len(rows)*155
    im=Image.new('RGB',(w,h),'white');dr=ImageDraw.Draw(im)
    dr.text((45,30),title,font=F,fill='#162539')
    for j,row in enumerate(rows):
        y=110+j*155; n=len(row);bw=(w-100-(n-1)*50)//n
        for k,s in enumerate(row):
            x=50+k*(bw+50);dr.rounded_rectangle((x,y,x+bw,y+95),radius=12,fill='#edf3f8',outline='#496782',width=3)
            lines=s.split('\n')
            for l,t in enumerate(lines):
                box=dr.textbbox((0,0),t,font=FS);dr.text((x+(bw-box[2])/2,y+18+l*30),t,font=FS,fill='#142d42')
            if k<n-1:
                dr.line((x+bw+5,y+47,x+bw+42,y+47),fill='#38566e',width=4);dr.polygon([(x+bw+42,y+47),(x+bw+31,y+40),(x+bw+31,y+54)],fill='#38566e')
    im.save(R/(name+'.png'))
diagram('architecture','System modules and event flow',[
    ['Input System','PlayerController\nPlayer FSM','Movement / Health / Hide'],
    ['Footstep / Rock / Trap','KillerAI\nEnemy FSM','Animator / UI danger'],
    ['ItemPickup','InventoryManager','ExitDoor FSM'],
    ['LevelConfig','GameManager\nScene management','HUD / Audio']])
diagram('player_fsm','Player FSM selected transitions',[
    ['Normal','Hide / Reveal','Hidden / Normal'],['Normal','Damage / Recover','Hurt / Normal'],['Normal','Freeze / Unfreeze','Frozen / Normal'],['Normal','Inventory open / close','UsingItem / Normal'],['Normal','Exit complete','Transitioning'],['Alive state','Health reaches zero','Dead']])
diagram('enemy_fsm','Enemy FSM main paths summary',[
    ['Patrol','Sound event','Listening'],['InvestigateSound','ConfirmSound','Search / ReturnToPatrol'],
    ['Player detected','Chase','Attack / AttackRecover'],['Lost target','Investigate / Search','ReturnToPatrol'],
    ['Scream ability','Scream','Player Stun / Frozen']])
diagram('warning_flow','TensionDirector main transitions and outputs',[
    ['Calm','Suspicious event','Suspicious'],['Any state','Killer hunting / Scream','Hunted'],['Hunted','Hunting ends','Escaped'],['Escaped','1.5 seconds','Cooldown'],['Cooldown','No suspicion / 2.2 sec','Calm'],['Tension state event','LevelAudioController','Chase music']])
diagram('raycast','Sight checks in KillerAI DetectPlayer',[
    ['Player hidden?','Hidden: no detection'],['Killer not blind','Within effective sight range'],
    ['Flashlight ON: any angle','Otherwise: within sight angle'],['Physics2D.Raycast\nfrom eye toward player','obstacleMask / distance'],['No obstacle collider','Sight detected'],['Sight not detected','Evaluate hearing separately']])

def montage(name,names):
    imgs=[Image.open(R/(s+'.png')).convert('RGB') for s in names]
    width=1400;gap=16;hh=0
    scaled=[]
    for im,s in zip(imgs,names):
        h=round(im.height*width/im.width);im=im.resize((width,h));scaled.append((im,s));hh+=h+52+gap
    out=Image.new('RGB',(width,hh-gap),'white');dr=ImageDraw.Draw(out);y=0
    for im,s in scaled:
        dr.text((15,y+8),s.replace('_',' ').upper(),font=FS,fill='black');y+=52;out.paste(im,(0,y));y+=im.height+gap
    out.save(R/(name+'.png'));return name
for name,names in [('movement',['walk','sprint','jump']),('damage',['hurt','gameover']),('endings',['complete','won']),('losswin',['gameover','won']),('states',['normal','hidden']),('levels',['level1','level2','level3']),('dangercompare',['level3','warning']),('keycompare',['pickup','collected'])]:montage(name,names)
fig3={1:'menu',2:'settings',3:'level1',4:'sprint',5:'pickup',6:'inventory',7:'hidden',8:'flashlight',9:'noise',10:'chase',11:'damage',12:'unlock',13:'endings',14:'curse',15:'level3',16:'trap',17:'warning',18:'architecture',19:'player_fsm',20:'enemy_fsm',21:'warning_flow'}
fig4={1:'movement',2:'flashlight',3:'states',4:'patrol',5:'noise',6:'chase',7:'raycast',8:'trap',9:'keycompare',10:'unlock',11:'level1',12:'losswin',13:'levels',14:'level3',15:'dangercompare'}
text(343,P[343].text)
text(350,P[350].text+' แผนภาพย่อแสดงคู่เหตุการณ์หลัก ไม่ได้แสดง Transition ทุกเส้น')
for i,p in enumerate(P):
    if (265<=i<359 or 439<=i<493) and (p.text.startswith('[แทรกภาพ') or (not p.text.strip() and p._p.xpath('.//w:drawing'))):remove(p)
for p in list(D.paragraphs):
    m=re.match(r'ภาพที่ ([34])-(\d+) ',p.text)
    if not m:continue
    chap,num=map(int,m.groups());name=(fig3 if chap==3 else fig4).get(num)
    if not name:continue
    pic=before(p);pic.alignment=WD_ALIGN_PARAGRAPH.CENTER;pic.paragraph_format.keep_with_next=True
    im=Image.open(R/(name+'.png'));width=min(5.9,5.25*im.width/im.height)
    pic.add_run().add_picture(str(R/(name+'.png')),width=Inches(width))
    # Each runtime figure is explicitly marked as controlled demonstration.
    if name not in ['architecture','player_fsm','enemy_fsm','warning_flow','raycast'] and 'สาธิต' not in p.text:p.text+=' (ภาพสาธิตใน Play Mode)'

# Replace every results placeholder with an honest evidence status.
tables=list(D.tables)
tables[0].rows[4].cells[2].text='โต้ตอบกับวัตถุ หรือกดค้างเพื่อปลดล็อกประตู; กุญแจเก็บเมื่อสัมผัส'
for t in tables[4:8]:
    for row in t.rows[1:]:row.cells[-1].text='แผนทดสอบ'
observed={
 'TC-01':'มีภาพเดินจากการส่งปุ่ม D; ยังไม่ตรวจครบทุกทิศ',
 'TC-02':'มีภาพจากการส่ง D กับ Shift; ยังไม่วัดความเร็วและเสียง',
 'TC-03':'มีภาพจากปุ่ม Space; ยังไม่ทดสอบกระโดดซ้ำ',
 'TC-05':'มีภาพไฟฉายหลังเรียก SetOn; ยังไม่ทดสอบการหมดแบตเตอรี่',
 'TC-FSM-P01':'เรียก ToggleHide หลังจัดเงื่อนไข; log เป็น Hidden และเห็น HIDDEN',
 'TC-FSM-P02':'เรียก ForceReveal; log กลับ Normal',
 'TC-FSM-P03':'เรียก TakeDamage(10); log เป็น Hurt',
 'TC-FSM-P04':'เรียก Freeze และ Unfreeze; log เป็น Frozen และกลับ Normal',
 'TC-FSM-P05':'เรียก OpenRoutine; log เป็น UsingItem และมีภาพคลัง',
 'TC-FSM-P06':'เรียก CompleteUnlock; log เป็น Transitioning และ LevelComplete',
 'TC-FSM-P07':'เรียก Kill; log เป็น Dead และมี Game Over; ไม่ใช่โจมตีต่อเนื่อง',
 'TC-FSM-E01':'ส่ง HearNoise ประเภท Rock; log จาก Patrol เป็น Listening',
 'TC-FSM-E04':'เรียก SetState(Chase); มี DANGER; ยังไม่ทดสอบการตรวจพบตามธรรมชาติ',
 'TC-06':'เรียก Collect; จำนวนกุญแจและภาพคลังเปลี่ยน',
 'TC-07':'มีภาพประตูก่อนเก็บครบ; ยังไม่ยืนยัน Prompt ทุกเงื่อนไข',
 'TC-08':'กำหนด Unlocking และความคืบหน้า; มี Prompt และเรียก CompleteUnlock',
 'TC-10':'จัดผู้เล่นที่กับดัก; log Killer เป็น Listening; ยังไม่ฟังเสียงครบลำดับ',
 'TC-12':'ภาพมีขอบจอสีแดงเมื่อจัดระยะใกล้; ยังไม่วัดเสียงหัวใจ',
 'TC-13':'เรียก NotifyLevelCompleted และ ContinueToNextLevel; เห็น YOU ESCAPED'}
for t in tables[8:12]:
    for row in t.rows[1:]:
        code=row.cells[0].text
        row.cells[3].text=observed.get(code,'ตรวจโค้ดพบเงื่อนไข; ยังไม่มีผลทดลองครบกรณี')
        row.cells[4].text='สาธิตบางส่วน' if code in observed else 'ยังไม่ทดสอบ'
for row in tables[12].rows[1:]:row.cells[1].text='ยังไม่มีข้อมูล';row.cells[2].text='ยังไม่ประเมิน'

# Move existing captions directly before their tables and add missing table labels.
for j in range(5):
    cap=next(p for p in D.paragraphs if p.text.startswith(f'ตารางที่ 4-{j+1} '))
    tables[8+j]._tbl.addprevious(cap._p)
table_titles=['คำสั่งควบคุมเกม','สถานะของ Player FSM','สถานะของ Enemy FSM','พารามิเตอร์ของแต่ละด่าน','แผนทดสอบการควบคุมพื้นฐาน','แผนทดสอบ Player FSM','แผนทดสอบ Enemy FSM','แผนทดสอบไอเท็มและ Level 3']
from docx.text.paragraph import Paragraph
for j,title in enumerate(table_titles):
    node=OxmlElement('w:p');tables[j]._tbl.addprevious(node);p=Paragraph(node,tables[j]._parent);p.add_run(f'ตารางที่ 3-{j+1} {title}')

# Add bibliography and appendices containing reproducible evidence and a mock calculation.
add('บรรณานุกรม','Heading 1')
refs=[
 'Jagdale, D. (2021). Finite State Machine in Game Development. International Journal of Advanced Research in Science, Communication and Technology, 10(1), 384–390. https://doi.org/10.48175/IJARSCT-2062',
 'Kusuma, A. F. A. A., & Handayani, R. D. (2024). Implementation of Inventory Checking System Using Finite State Machine in Role Playing Game. International Journal of Computer and Information System, 5(4). https://doi.org/10.29040/ijcis.v5i4.197',
 'Thyrun, V. (2026). An Analysis of Finite State Machine Based Enemy Artificial Intelligence in Kirby: Nightmare in Dream Land. Preprints.org [บทความก่อนประเมินโดยผู้ทรงคุณวุฒิ]. https://doi.org/10.20944/preprints202601.1473.v1',
 'Koulaxidis, G., & Xinogalos, S. (2022). Improving Mobile Game Performance with Basic Optimization Techniques in Unity. Modelling, 3, 201–223. https://doi.org/10.3390/modelling3020014',
 'Unity Technologies. (2026). Animator Controller. Unity Manual 6.4. สืบค้น 4 ตุลาคม 2569 จาก https://docs.unity3d.com/6000.4/Documentation/Manual/class-AnimatorController.html',
 'Unity Technologies. (2026). Animation Parameters. Unity Manual 6.4. สืบค้น 4 ตุลาคม 2569 จาก https://docs.unity3d.com/6000.4/Documentation/Manual/AnimationParameters.html',
 'Unity Technologies. (2026). Physics2D.Raycast. Unity Scripting API 6.4. สืบค้น 4 ตุลาคม 2569 จาก https://docs.unity3d.com/6000.4/Documentation/ScriptReference/Physics2D.Raycast.html',
 'ผู้จัดทำ. (2569). SideScrollingEscape [โค้ดต้นแบบและทรัพยากรเกม Unity]. PlayerController.cs, KillerAI.cs, PlayerAnimator.controller, KillerAnimator.controller, InventoryManager.cs, ExitDoor.cs, UIManager.cs และ LevelConfig. ตรวจสอบ 4 ตุลาคม 2569.']
for s in refs:add(s)
add('ภาคผนวก ก หลักฐานและวิธีการสาธิต','Heading 1')
add('ก.1 แหล่งภาพและข้อจำกัด','Heading 2')
add('ภาพเกมในบท 3 และ 4 มาจาก Unity Play Mode ของโครงการ ภาพ Animator ที่ 3-22 และ 3-23 มาจากหน้าต่าง Animator ของ Unity ส่วน 4-16–4-21 เป็นเฟรมจาก Sprite Sheet จริง แผนภาพที่ 3-18–3-21 และ 4-7 วาดสรุปจากโค้ดเพื่ออธิบายโครงสร้าง ไม่ใช่ภาพจับจากการเล่น')
add('สคริปต์สาธิตชั่วคราวใช้ส่งปุ่ม D, Shift และ Space เรียก SetOn, Collect, OpenRoutine, ToggleHide, ForceReveal, Freeze, Unfreeze, TakeDamage, HearNoise, SetState, CompleteUnlock, Kill และ ApplyStatueCurse รวมถึงคำสั่งเปลี่ยนด่าน บางภาพจัดตำแหน่งผู้เล่นและ Killer ปิดการอัปเดตศัตรู ปิดโอกาสสุ่ม Jumpscare หรือคงสถานะประตูเพื่อแสดงผล ภาพดังกล่าวไม่ใช้ยืนยันว่าผู้เล่นทำสำเร็จผ่านการเล่นตามปกติ')
add('ก.2 บันทึกสถานะที่สังเกตได้','Heading 2')
add('ตารางที่ ก-1 ตัวอย่างสถานะจากบันทึกการสาธิต')
t=D.add_table(rows=1,cols=3);t.rows[0].cells[0].text='เหตุการณ์สาธิต';t.rows[0].cells[1].text='สถานะผู้เล่น';t.rows[0].cells[2].text='ผลแสดงหรือสถานะเกม'
for row in [('เปิดคลัง','UsingItem','หน้าต่าง Inventory'),('เข้าและออกจากการซ่อน','Hidden → Normal','HIDDEN แสดงและหาย'),('หยุดผู้เล่น','Frozen','จัดสถานะหยุดชั่วคราว'),('ความเสียหาย 10 หน่วย','Hurt','แถบ Health เปลี่ยน'),('ส่งเสียง Rock','Normal','Killer: Listening'),('เรียก Chase','Normal','Killer: Chase / DANGER'),('เรียก CompleteUnlock','Transitioning','LevelComplete'),('เรียก Kill','Dead','Dead / Game Over'),('ผ่านเงื่อนไขด่านสุดท้ายโดยคำสั่ง','Normal','GameWon / YOU ESCAPED')]:
    cells=t.add_row().cells
    for c,s in zip(cells,row):c.text=s
add('ก.3 วิธีทดสอบต่อ','Heading 2')
add('คืนการทำงานตามปกติของศัตรูและการสุ่มเหตุการณ์ เล่นแต่ละ Test Case ด้วยปุ่มควบคุมจริง บันทึกผลเทียบเกณฑ์ทีละข้อ แล้วเล่นจากเมนูจนผ่านทุกด่านอย่างต่อเนื่อง ควรตรวจ Console ปัญหาการชน แสง กล้อง เสียง และการยกเลิกปลดล็อก จากนั้นจึงเก็บแบบสอบถามกับผู้ทดลองเล่น')
add('ภาคผนวก ข ตัวอย่างคะแนนสมมติและแบบประเมิน','Heading 1')
add('ข.1 ข้อมูลสมมติสำหรับแสดงวิธีคำนวณ','Heading 2')
add('ข้อมูลทั้งหมดในหัวข้อนี้สร้างขึ้นเพื่อเป็นตัวอย่าง ไม่ได้มาจากผู้ทดลองเล่นจริง รหัส S1–S5 หมายถึงแถวข้อมูลสมมติ 5 แถว ไม่ใช่จำนวนกลุ่มตัวอย่างจริง จึงห้ามนำค่าเฉลี่ยหรือส่วนเบี่ยงเบนมาตรฐานนี้ไปอ้างว่าเป็นผลความพึงพอใจของเกม')
add('ตารางที่ ข-1 คะแนนสมมติสำหรับตัวอย่างการคำนวณเท่านั้น')
mock=[[4,4,4,3,5,4],[5,4,4,4,4,4],[4,5,3,4,5,5],[4,4,4,3,4,4],[5,4,5,4,4,5]]
t=D.add_table(rows=1,cols=7)
for c,s in zip(t.rows[0].cells,['รหัสสมมติ','ควบคุม','UI','ศัตรู','ความยาก','บรรยากาศ','โดยรวม']):c.text=s
for i,row in enumerate(mock):
    for c,s in zip(t.add_row().cells,[f'S{i+1}']+row):c.text=str(s)
add('ค่าเฉลี่ย x̄ = Σxi / n และส่วนเบี่ยงเบนมาตรฐานของตัวอย่าง s = √[Σ(xi − x̄)² / (n − 1)] โดย n ในตัวอย่างนี้คือจำนวนแถวสมมติ สำหรับด้านควบคุม ผลรวมเท่ากับ 22 จึงได้ค่าเฉลี่ย 22/5 = 4.40 และ s = 0.55 เมื่อปัดสองตำแหน่ง')
add('ตารางที่ ข-2 ผลคำนวณจากคะแนนสมมติเท่านั้น')
t=D.add_table(rows=1,cols=4)
for c,s in zip(t.rows[0].cells,['ด้านสมมติ','ค่าเฉลี่ย','s','แปลผลตามเกณฑ์ตัวอย่าง']):c.text=s
for j,name in enumerate(['การควบคุม','UI','พฤติกรรมศัตรู','ความยาก','บรรยากาศ','ความพึงพอใจโดยรวม']):
    vals=[r[j] for r in mock];mean=statistics.mean(vals)
    for c,s in zip(t.add_row().cells,[name,f'{mean:.2f}',f'{statistics.stdev(vals):.2f}','ดีมาก' if mean>=4.21 else 'ดี']):c.text=s
add('เกณฑ์ตัวอย่างแบ่งค่าเฉลี่ยเป็น 4.21–5.00 ดีมาก, 3.41–4.20 ดี, 2.61–3.40 ปานกลาง, 1.81–2.60 น้อย และ 1.00–1.80 น้อยมาก ในข้อมูลสมมตินี้บรรยากาศกับความพึงพอใจโดยรวมมีค่าเฉลี่ย 4.40 และด้านความยาก 3.60 เป็นเพียงการฝึกแปลผล ยังไม่ใช่ข้อค้นพบจากเกมหรือผู้ใช้งานจริง')
add('ข.2 แบบประเมินสำหรับเก็บข้อมูลจริง','Heading 2')
add('ให้ผู้ตอบทดลองเล่นด้วยระบบตามปกติ แล้วเลือกคะแนนแต่ละข้อจาก 1 ถึง 5 โดย 1 น้อยมาก 2 น้อย 3 ปานกลาง 4 ดี และ 5 ดีมาก แบบฟอร์มนี้ยังไม่ได้รับคำตอบจริง')
add('วันที่ทดลอง __________________  อุปกรณ์ที่ใช้ __________________')
add('ประสบการณ์เล่นเกมสยองขวัญ: ไม่เคย / บางครั้ง / เป็นประจำ')
add('ตารางที่ ข-3 แบบประเมินสำหรับผู้ทดลองเล่นจริง')
t=D.add_table(rows=1,cols=6)
for c,s in zip(t.rows[0].cells,['รายการประเมิน','1','2','3','4','5']):c.text=s
for s in ['ควบคุมตัวละครได้สะดวก','UI และข้อความแนะนำชัดเจน','พฤติกรรมศัตรูเหมาะสม','ระดับความยากมีความสมดุล','บรรยากาศสอดคล้องกับเกมสยองขวัญ','ความพึงพอใจโดยรวม']:
    cells=t.add_row().cells;cells[0].text=s
    for c in cells[1:]:c.text='○'
add('คำถามปลายเปิด: ส่วนใดเข้าใจยากที่สุด เหตุการณ์ใดทำให้รู้สึกไม่ยุติธรรม พบปัญหาหรือข้อผิดพลาดใด และต้องการให้ปรับปรุงอะไร ควรเก็บคำตอบดิบไว้ตรวจสอบและแยกข้อเสนอแนะจริงออกจากความคิดเห็นของผู้พัฒนา')
add('ข้อเสนอแนะ __________________________________________________')
add('____________________________________________________________')
add('____________________________________________________________')

# Rebuild the original metadata block as aligned labels and values.
meta=D.add_table(rows=0,cols=2)
for label,value in [('หัวข้อโครงงาน','การพัฒนาเกมโดยประยุกต์ใช้เทคนิค Finite State Machine ในการควบคุมตัวละครและปัญญาประดิษฐ์\nDevelopment of a Game using Finite State Machine Techniques for Character Control and Artificial Intelligence'),('ผู้จัดทำ','นายจิรายุ สงบวาจา\nนางสาวอรไพลิน สอาดเอี่ยม\nนางสาววนิดา คล้ายสกุล'),('อาจารย์ที่ปรึกษา','อาจารย์ธนวัฒน์ ถาวรกูล\nผศ.วราภรณ์ มั่นทุ่ง'),('ระดับการศึกษา','ปริญญาตรี'),('สาขาวิชา','วิทยาการคอมพิวเตอร์'),('ปีการศึกษา','2569')]:
    cells=meta.add_row().cells;cells[0].text=label;cells[1].text=value
P[33]._p.addprevious(meta._tbl)
for p in P[33:46]:remove(p)

# Structure, typography, pagination, lists and page-number fields.
for name in ['FigureCaption','TableCaption']:
    if name not in D.styles:D.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
for s in D.styles:
    if s.type==WD_STYLE_TYPE.PARAGRAPH and s.name in ['Normal','Title','Heading 1','Heading 2','Heading 3','FigureCaption','TableCaption','TOC 1','TOC 2']:
        s.font.name='TH SarabunPSK';s.font.size=Pt(16);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.line_spacing=1;s.paragraph_format.space_after=Pt(5)
for p in list(D.paragraphs):
    # Remove empty layout filler outside covers, while keeping inline art.
    if not p.text.strip() and not p._p.xpath('.//w:drawing'):
        if p not in P[:33]:remove(p)
for p in D.paragraphs:
    s=p.text.strip();f=p.paragraph_format;f.line_spacing=1;f.space_after=Pt(5);f.widow_control=True;f.page_break_before=False
    if p.text!=s and not p._p.xpath('.//w:drawing'):p.text=s
    if s.startswith('ภาพที่ '):p.style='FigureCaption';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.first_line_indent=Pt(0);f.keep_with_next=False
    elif s.startswith('ตารางที่ '):p.style='TableCaption';f.first_line_indent=Pt(0);f.keep_with_next=True
    elif s.startswith('บทที่ ') or s in ['บรรณานุกรม'] or s.startswith('ภาคผนวก '):p.style='Heading 1';f.page_break_before=True;f.keep_with_next=True;f.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif re.match(r'^[1-5]\.\d+\s',s):p.style='Heading 2';f.keep_with_next=True;f.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    elif re.match(r'^[1-5]\.\d+\.\d+\s',s) or re.match(r'^[กข]\.\d+\s',s):p.style='Heading 3';f.keep_with_next=True;f.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    elif s and p not in P[:46]:
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;f.first_line_indent=Inches(.5);f.keep_with_next=False
    for run in p.runs:font(run,18 if p.style.name.startswith('Heading') else 16,p.style.name.startswith('Heading'))
for i in [14,64,359,367,374,386]:P[i].paragraph_format.page_break_before=True
for p in P[:33]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
for p in P[151:161]:
    p.style='Normal';p.paragraph_format.keep_with_next=False
    for run in p.runs:font(run,16,False)
for i in [1,14]:P[i].style='Title';P[i].paragraph_format.keep_with_next=True
for i in [46,64]:P[i].alignment=WD_ALIGN_PARAGRAPH.CENTER;P[i].paragraph_format.keep_with_next=True
for p in P[69:73]:
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Pt(0);p.paragraph_format.right_indent=Pt(0)
# Chapter names join the chapter heading in the TOC without creating separate entries.
for i,j in [(81,82),(148,149),(216,217),(437,438),(574,575)]:
    P[i].text=P[i].text+' '+P[j].text;remove(P[j]);font(P[i].runs[0],18,True)
for t in D.tables:
    t.autofit=False;t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    n=len(t.columns); widths={2:[1.4,4.5],3:[2.1,1.5,2.3],4:[2.4,1,1,1.5],5:[.65,1.1,1.35,1.85,.95],6:[3.4,.5,.5,.5,.5,.5],7:[.8,.85,.85,.85,.85,.85,.85]}.get(n,[5.9/n]*n)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr
    for old in pr.findall(qn('w:tblInd')):pr.remove(old)
    for old in pr.findall(qn('w:tblpPr')):pr.remove(old)
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:w'),str(round(sum(widths)*1440)));tw.set(qn('w:type'),'dxa')
    for row in t.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for c,w in zip(row.cells,widths):
            c.width=Inches(w)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=False
                for run in p.runs:font(run,14,row==t.rows[0])
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for c in t.rows[0].cells:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');c._tc.get_or_add_tcPr().append(sh)
        for p in c.paragraphs:
            for run in p.runs:font(run,14,True)
meta.style='Normal Table'
for row in meta.rows:
    for j,c in enumerate(row.cells):
        for sh in c._tc.get_or_add_tcPr().findall(qn('w:shd')):c._tc.get_or_add_tcPr().remove(sh)
        for p in c.paragraphs:
            for run in p.runs:font(run,16,j==0)
meta.rows[0].cells[0].paragraphs[0].paragraph_format.page_break_before=True
for p in D.paragraphs:
    if p.text.startswith('ข.2 '):p.paragraph_format.page_break_before=True
    if p.text.startswith('ตารางที่ 4-1 '):p.paragraph_format.page_break_before=True
for t in D.tables:
    for row in t.rows:
        for c in row.cells:
            if c.text=='แปลผลตามเกณฑ์ตัวอย่าง':c.text='ระดับตัวอย่าง';font(c.paragraphs[0].runs[0],14,True)
for section in D.sections:
    footer=section.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(footer,'PAGE');section.different_first_page_header_footer=True
# Native Word fields create and maintain page references.
for title,instr in [('สารบัญ','TOC \\o "1-2" \\h \\z'),('สารบัญภาพ','TOC \\t "FigureCaption,1" \\h \\z'),('สารบัญตาราง','TOC \\t "TableCaption,1" \\h \\z')]:
    h=before(P[81],title);h.paragraph_format.page_break_before=True;h.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(h.runs[0],18,True)
    p=before(P[81]);field(p,instr)
settings=D.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
D.save(OUT)
print('Saved',OUT.name,'paragraphs',len(D.paragraphs),'tables',len(D.tables),'pictures',len(D.inline_shapes))
assert not any('กรอก' in p.text or '[แทรกภาพ' in p.text for p in D.paragraphs)
assert not any('[กรอก' in c.text for t in D.tables for row in t.rows for c in row.cells)
