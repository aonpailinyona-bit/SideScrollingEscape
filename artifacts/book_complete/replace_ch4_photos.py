from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

r=Path(__file__).parent
src=Path('D:/งาน/project_บท3_ภาพออกแบบครบ20รูป.docx')
out=r/'project_บท4_ปรับภาพครบ17รูป.docx'
d=Document(src)
ps=d.paragraphs
a=next(i for i,p in enumerate(ps) if p.text.startswith('บทที่ 4'))
b=next(i for i,p in enumerate(ps[a+1:],a+1) if p.text=='บทที่ 5')
before=[p._p.xml for p in ps[:a]];after=[p._p.xml for p in ps[b:]]
tables=[t._tbl.xml for t in d.tables]
for p in ps[a:b]:
 if p._p.xpath('.//w:drawing') or (p.text.startswith('ภาพที่ 4-') and not p.text.startswith('ภาพที่ 4-16 และ 4-17')):
  p._p.getparent().remove(p._p)
def replace(prefix,text):
 p=next(p for p in d.paragraphs[a:] if p.text.startswith(prefix))
 p.text=text
 for run in p.runs:run.font.name='TH SarabunPSK';run.font.size=Pt(16)
replace('บทนี้นำเสนอ','บทนี้นำเสนอผลการพัฒนาต้นแบบเกมสามด่าน ภาพประกอบเป็นภาพหน้าจอที่จัดเก็บจากเกมและแสดงผลการทำงานที่มองเห็นได้ เช่น การเคลื่อนที่ ไฟฉาย คลังไอเท็ม การซ่อนตัว คำสาป เมนู และหน้าจอผลลัพธ์ การตรวจโครงสร้างและผลทดสอบองค์ประกอบรายงานแยกในหัวข้อ 4.9 และ 4.10 โดยผลทดสอบ 35 กรณีครอบคลุมเงื่อนไขควบคุมที่กำหนด ส่วนการเล่นครบทุกเส้นทางและการประเมินจากผู้ใช้งานจริงยังต้องดำเนินการเพิ่มเติม')
replace('บันทึกการสาธิตพบ','ภาพหน้าจอแสดงข้อความ HIDDEN เมื่อผู้เล่นซ่อนตัว และข้อความ PETRIFIED — BREAK FREE! เมื่อผู้เล่นอยู่ในสถานะถูกทำให้หยุดเคลื่อนไหว โดยภาพใน Level 2 แสดงตัวนับและคำแนะนำให้กดปุ่ม X เพื่อหลุดจากสถานะ ภาพเหล่านี้แสดงผลที่ผู้เล่นมองเห็น ส่วนการเปลี่ยนสถานะของ Player FSM อ้างอิงผลทดสอบในหัวข้อ 4.10')
replace('TensionDirector รวม','TensionDirector รวมสถานะของศัตรูเป็น Calm, Suspicious, Hunted, Escaped และ Cooldown แล้วส่งเหตุการณ์ให้ LevelAudioController ปรับดนตรี ส่วน UIManager แสดง DANGER และ Level3HeartbeatWarning ควบคุมคำเตือนตามระยะห่าง ภาพหน้าจอ Level 2 แสดง DANGER ขอบจอสีแดง และข้อความ CURSE RISING — BREAK LINE OF SIGHT! ซึ่งช่วยสื่อสารสถานการณ์อันตรายและคำแนะนำในการหลบหลีก ภาพนี้ไม่ได้แสดงหน้าต่าง Debug ของ TensionDirector จึงใช้ประกอบการอธิบายการแจ้งเตือนบนหน้าจอ')
replace('สิ่งที่เห็นจากภาพ','สิ่งที่เห็นจากภาพคือข้อความเตือน แถบคำสาป และขอบจอสีแดง ส่วนความชัดเจนของเสียงและความเหมาะสมของระดับความกดดันต้องตรวจด้วยการฟังและทดลองเล่นเพิ่มเติม')
replace('ภาพที่ 4-16 และ 4-17','คลิป PlayerIdle.anim และ PlayerWalk.anim ใช้สำหรับท่ายืนและท่าเดินตามลำดับ ทั้งสองคลิปกำหนด Sample Rate เท่ากับ 12 และเปิด Loop Time เพื่อเล่นวนขณะตัวละครคงอยู่ในสถานะเดิม โดยคลิป Idle มีช่วงเวลา 0.75 วินาที ส่วน Walk มีช่วงเวลา 1 วินาที แผนผัง Controller ของผู้เล่นแสดงในภาพที่ 3-25')
replace('KillerAnimator มีสถานะ','KillerAnimator มีสถานะ Idle, Walk, Run, Attack1 และ Scream โดยเชื่อมกับไฟล์ KillerIdle.anim, KillerWalk.anim, KillerRun.anim, KillerAttack1.anim และ KillerScream.anim เพื่อแสดงการหยุดนิ่ง การเดิน การวิ่ง การโจมตี และการกรีดร้อง แผนผัง Controller และ Parameters แสดงในภาพที่ 3-26 และผลทดสอบการเปลี่ยนสถานะรายงานในตารางที่ 4-8')
for p in d.paragraphs[a:]:
 if p.text.startswith('จากการตรวจสอบโครงสร้าง PlayerAnimator'):
  p.text=p.text.replace('ส่วนความถูกต้องของภาพที่แสดงเมื่อกดหรือปล่อยปุ่มยังต้องตรวจใน Play Mode','ผลทดสอบ Controller ใน Play Mode ตาม AN-P01 ผ่านตามตารางที่ 4-8')
  for run in p.runs:run.font.name='TH SarabunPSK';run.font.size=Pt(16)
groups=[
('4.2 ',[(15,'ตัวละครผู้เล่นระหว่างเคลื่อนที่ใน Level 3','ภาพแสดงตัวละครบนเส้นทางบริเวณทางเข้าสวนสนุก โดย HUD แสดงด่าน จำนวนกุญแจ และทรัพยากรที่เหลือระหว่างการเคลื่อนที่'),(14,'ตัวละครผู้เล่นขณะกระโดดใน Level 3','ภาพแสดงตัวละครยกตัวเหนือพื้นบริเวณช่องว่างของเส้นทาง สื่อให้เห็นการเคลื่อนที่กลางอากาศขณะกระโดด'),(1,'การใช้ไฟฉายใน Level 1','ภาพแสดงลำแสงไฟฉายด้านหน้าตัวละครและกุญแจบนเส้นทาง พร้อมแถบ Battery สำหรับติดตามแบตเตอรี่คงเหลือ')]),
('4.3 ',[(7,'การแสดงสถานะ Hidden ใน Level 2','เมื่อผู้เล่นซ่อนตัว หน้าจอแสดง HIDDEN และข้อความ Press E to Exit Hiding เพื่อแจ้งสถานะและวิธีออกจากจุดซ่อน'),(8,'การแสดงสถานะ Petrified ใน Level 2','ภาพแสดงคำเตือน PETRIFIED — BREAK FREE! พร้อมตัวนับและคำแนะนำให้กด X เพื่อหลุดจากสถานะ ขณะ HUD ยังคงแสดงข้อมูลพลังชีวิตและคำสาป')]),
('4.6 ',[(5,'คลังไอเท็มและรายการสิ่งของที่เก็บได้','ภาพแสดงหน้าต่าง Inventory ที่มีช่องเก็บกุญแจและขวดไอเท็ม พร้อมจำนวนกุญแจบน HUD ทำให้ผู้เล่นตรวจสอบทรัพยากรได้ระหว่างเล่น'),(10,'ประตูทางออกและคำสั่งปลดล็อกใน Level 2','เมื่อจำนวนกุญแจบน HUD ครบ 5/5 หน้าจอแสดง Hold E to UNLOCK บริเวณประตู เพื่อแนะนำให้ผู้เล่นกด E ค้างเพื่อปลดล็อก'),(0,'หน้าจอ Level Complete','เมื่อผ่านด่าน หน้าจอแสดง LEVEL COMPLETE พร้อมปุ่ม Next Level และ Main Menu ให้เลือกเล่นด่านถัดไปหรือกลับเมนูหลัก')]),
('4.7 ',[('image','หน้าจอแนะนำวิธีเล่นของ Level 1','หน้าจอแนะนำแสดงชื่อด่าน THE FORSAKEN CELLAR และปุ่มควบคุม เช่น A/D สำหรับเดิน Shift สำหรับวิ่ง Space สำหรับกระโดด E สำหรับโต้ตอบ และปุ่มถัดไป'),(2,'หน้าจอหยุดพักเกม','หน้าจอ PAUSED แสดงตัวเลือก Resume, Settings, Main Menu และ Exit Game เพื่อให้ผู้เล่นเลือกดำเนินการระหว่างหยุดพัก'),(3,'หน้าจอตั้งค่าเสียง','หน้าจอ SETTINGS แสดงแถบเลื่อน Master, SFX และ Music พร้อมปุ่ม Back เพื่อปรับระดับเสียงและย้อนกลับ'),(4,'หน้าจอเมื่อตัวละครเสียชีวิต','หน้าจอ YOU DIED แสดงปุ่ม Try Again และ Main Menu เพื่อให้ผู้เล่นเลือกเริ่มใหม่หรือกลับเมนูหลัก'),(16,'หน้าจอเมื่อหลบหนีสำเร็จ','หน้าจอ YOU ESCAPED แสดงเมื่อจบเกมสำเร็จ พร้อมปุ่ม Main Menu เพื่อกลับเข้าสู่เมนูหลัก')]),
('4.8 ',[(11,'พื้นที่ทางเข้าสวนสนุกใน Level 3','ภาพแสดงพื้นที่ทางเข้าสวนสนุก ป้าย CIRCUS ทางเดิน กุญแจ และองค์ประกอบฉาก พร้อมชื่อด่าน Echoes of the Blind ระดับ Nightmare บน HUD'),(12,'พื้นที่เต็นท์และจุดซ่อนใน Level 3','ภาพแสดงเต็นท์ขนาดใหญ่ ถังที่ใช้เป็นจุดซ่อน และข้อความ Press E to Hide เพื่อสื่อสารตำแหน่งวัตถุที่โต้ตอบได้'),(13,'เส้นทางบริเวณปืนใหญ่ข้างเต็นท์ใน Level 3','ภาพแสดงตัวละครบริเวณปืนใหญ่ด้านข้างเต็นท์และทางเดินของสวนสนุก ซึ่งเป็นองค์ประกอบฉากที่ช่วยให้ผู้เล่นจดจำตำแหน่งระหว่างสำรวจ')]),
('4.9 ',[(6,'คำเตือน Danger และคำสาปใน Level 2','ภาพแสดง DANGER ขอบจอสีแดง และข้อความ CURSE RISING — BREAK LINE OF SIGHT! เพื่อแจ้งให้ผู้เล่นหลบจากแนวสายตา พร้อมแถบ CURSE สำหรับติดตามสถานะคำสาป')]),
]
num=0
for prefix,items in groups:
 anchor=next(p for p in d.paragraphs if p.text.startswith(prefix))
 for file,title,body in items:
  num+=1
  p=anchor.insert_paragraph_before(body);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.paragraph_format.first_line_indent=Inches(.35)
  p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(6)
  for run in p.runs:run.font.name='TH SarabunPSK';run.font.size=Pt(16)
  p=anchor.insert_paragraph_before();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
  p.add_run().add_picture(str(Path('D:/รูป')/f'{file}.png'),width=Inches(5.65))
  p=anchor.insert_paragraph_before(f'ภาพที่ 4-{num} {title}');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8)
  for run in p.runs:run.font.name='TH SarabunPSK';run.font.size=Pt(14)
# Remove empty paragraphs left by former figure spacing in chapter 4 only.
active=False
for p in list(d.paragraphs):
 if p.text.startswith('บทที่ 4'):active=True
 if p.text=='บทที่ 5':active=False
 if active and not p.text.strip() and not p._p.xpath('.//w:drawing') and not p._p.xpath('.//w:sectPr'):
  p._p.getparent().remove(p._p)
d.save(out)
check=Document(out);ps=check.paragraphs;ca=next(i for i,p in enumerate(ps) if p.text.startswith('บทที่ 4'));cb=next(i for i,p in enumerate(ps[ca+1:],ca+1) if p.text=='บทที่ 5')
assert before==[p._p.xml for p in ps[:ca]]
assert after==[p._p.xml for p in ps[cb:]]
assert tables==[t._tbl.xml for t in check.tables]
assert sum(len(p._p.xpath('.//w:drawing')) for p in ps[ca:cb])==17
assert num==17
print(out);print('PASS 17 new figures; chapters outside 4 and all tables preserved')
