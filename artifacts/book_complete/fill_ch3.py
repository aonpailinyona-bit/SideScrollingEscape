from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
import re
r=Path(__file__).resolve().parent
d=Document(r/'project_แก้เฉพาะบทที่4.docx')
b=list(d.element.body)
txt=lambda e: ''.join(e.xpath('.//w:t/text()'))
s=next(i for i,e in enumerate(b) if txt(e)=='บทที่ 3')
z=next(i for i,e in enumerate(b) if i>s and txt(e).startswith('บทที่ 4'))
before=[etree.tostring(e) for e in b[:s]];after=[etree.tostring(e) for e in b[z:]]
mapping={3:'level1',4:'sprint',5:'pickup',6:'inventory',7:'hidden',8:'flashlight',9:'noise',10:'chase',11:'damage',12:'unlock',13:'endings',14:'curse',15:'level3',16:'trap',17:'warning'}
from docx.text.paragraph import Paragraph
for e in b[s:z]:
    m=re.match(r'ภาพที่ 3-(\d+) ',txt(e))
    if not m or int(m[1]) not in mapping:continue
    n=int(m[1]);p=Paragraph(e,d._body)
    prev=e.getprevious()
    if prev is not None and '[แทรกภาพ' in txt(prev):prev.getparent().remove(prev)
    im=p.insert_paragraph_before();im.alignment=WD_ALIGN_PARAGRAPH.CENTER
    im.paragraph_format.keep_with_next=True
    im.add_run().add_picture(str(r/(mapping[n]+'.png')),width=Inches(5.65))
    labels={5:'สาธิตตำแหน่งกุญแจและระบบเก็บไอเท็ม',9:'สาธิต Killer รับเหตุการณ์เสียงประเภท Rock',15:'สภาพแวดล้อมสวนสนุกและพื้นที่ป่าใน Level 3',16:'สาธิตตำแหน่งกับดักเสียงใน Level 3'}
    if n in labels:p.text=f'ภาพที่ 3-{n} '+labels[n]
    p.add_run(' (ภาพสาธิตใน Play Mode)')
    p.paragraph_format.keep_with_next=False
    for run in p.runs:run.font.name='TH SarabunPSK';run.font.size=Pt(14)
out=r/'project_บท3เติมภาพ_บท4ล่าสุด.docx';d.save(out)
c=Document(out);bb=list(c.element.body);ss=next(i for i,e in enumerate(bb) if txt(e)=='บทที่ 3');zz=next(i for i,e in enumerate(bb) if i>ss and txt(e).startswith('บทที่ 4'))
assert before==[etree.tostring(e) for e in bb[:ss]]
assert after==[etree.tostring(e) for e in bb[zz:]]
assert not any('[แทรกภาพ' in txt(e) for e in bb[ss:zz])
print('Added 15 missing chapter 3 figures. All other chapters preserved.')
