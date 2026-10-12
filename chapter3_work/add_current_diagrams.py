from pathlib import Path
from math import atan2, cos, sin

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(r"C:\Users\aonpa\OneDrive\Desktop\SideScrollingEscape")
INPUT = ROOT / "project_บทที่3_ปรับปรุง.docx"
OUTPUT = ROOT / "project_บทที่3_ปรับปรุงพร้อมไดอะแกรม.docx"
ASSETS = ROOT / "chapter3_work" / "diagrams"
ASSETS.mkdir(parents=True, exist_ok=True)

FONT_REG = r"C:\Windows\Fonts\tahoma.ttf"
FONT_BOLD = r"C:\Windows\Fonts\tahomabd.ttf"
NAVY = "#24567F"
PALE = "#EAF2F8"
PALE_GREEN = "#E8F5EE"
PALE_RED = "#FBEDEE"
GRAY = "#6B7280"
LINE = "#748596"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def canvas(title, subtitle=""):
    im = Image.new("RGB", (1800, 1120), "white")
    d = ImageDraw.Draw(im)
    d.text((900, 50), title, font=font(44, True), fill="#111827", anchor="ma")
    if subtitle:
        d.text((900, 108), subtitle, font=font(24), fill=GRAY, anchor="ma")
    return im, d


def box(d, xy, title, detail="", fill=PALE, outline=NAVY, radius=22, title_size=27, detail_size=20):
    x1, y1, x2, y2 = xy
    d.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=4)
    cy = (y1 + y2) / 2
    if detail:
        d.text(((x1+x2)/2, cy-17), title, font=font(title_size, True), fill="#172033", anchor="mm")
        d.text(((x1+x2)/2, cy+25), detail, font=font(detail_size), fill="#354052", anchor="mm")
    else:
        d.text(((x1+x2)/2, cy), title, font=font(title_size, True), fill="#172033", anchor="mm")


def arrow(d, start, end, label="", dashed=False, color=LINE, width=4, label_offset=(0, -14)):
    x1, y1 = start; x2, y2 = end
    if dashed:
        steps = 18
        for i in range(0, steps, 2):
            a=i/steps; b=min((i+1)/steps,1)
            d.line((x1+(x2-x1)*a,y1+(y2-y1)*a,x1+(x2-x1)*b,y1+(y2-y1)*b), fill=color, width=width)
    else:
        d.line((x1,y1,x2,y2), fill=color, width=width)
    ang=atan2(y2-y1,x2-x1); size=16
    p1=(x2-size*cos(ang-0.55),y2-size*sin(ang-0.55)); p2=(x2-size*cos(ang+0.55),y2-size*sin(ang+0.55))
    d.polygon([(x2,y2),p1,p2], fill=color)
    if label:
        mx=(x1+x2)/2+label_offset[0]; my=(y1+y2)/2+label_offset[1]
        bbox=d.textbbox((mx,my),label,font=font(18),anchor="mm")
        d.rounded_rectangle((bbox[0]-7,bbox[1]-3,bbox[2]+7,bbox[3]+3),radius=6,fill="white")
        d.text((mx,my),label,font=font(18),fill="#303846",anchor="mm")


def save(im, name):
    path=ASSETS/name
    im.save(path, dpi=(220,220), optimize=True)
    return path


def architecture():
    im,d=canvas("แผนภาพสถาปัตยกรรมของระบบ", "โครงสร้างการทำงานตามสคริปต์ในโครงการเวอร์ชันล่าสุด")
    nodes={
        "input":(70,210,390,330), "player":(520,190,900,350), "foot":(1030,170,1390,290),
        "hide":(1030,320,1390,440), "inventory":(70,510,390,630), "exit":(520,500,900,640),
        "killer":(1030,500,1390,660), "tension":(1440,500,1740,660),
        "game":(520,790,900,930), "ui":(1030,790,1390,930), "audio":(1440,790,1740,930)
    }
    box(d,nodes["input"],"Input System","WASD, Shift, Space, E, F, H, C, Q")
    box(d,nodes["player"],"PlayerController","movement + Player FSM")
    box(d,nodes["foot"],"FootstepSystem","ส่งเหตุการณ์เสียง")
    box(d,nodes["hide"],"HidingSystem","จุดซ่อนและสถานะ Hidden")
    box(d,nodes["inventory"],"InventoryManager","จำนวนกุญแจและไอเท็ม")
    box(d,nodes["exit"],"ExitDoor","Locked ถึง Completed")
    box(d,nodes["killer"],"KillerAI","vision + hearing + Enemy FSM",fill=PALE_RED)
    box(d,nodes["tension"],"TensionDirector","ระดับความตึงเครียด",fill=PALE_RED)
    box(d,nodes["game"],"GameManager","สถานะเกมและเปลี่ยนฉาก",fill=PALE_GREEN)
    box(d,nodes["ui"],"UIManager","HUD, prompt และผลลัพธ์",fill=PALE_GREEN)
    box(d,nodes["audio"],"AudioManager","ดนตรีและเสียงเตือน",fill=PALE_GREEN)
    arrow(d,(390,270),(520,270),"คำสั่ง")
    arrow(d,(900,245),(1030,230),"การเคลื่อนไหว")
    arrow(d,(900,315),(1030,380),"เข้า/ออกจุดซ่อน")
    arrow(d,(1210,290),(1210,500),"เสียง")
    arrow(d,(390,570),(520,570),"จำนวนกุญแจ")
    arrow(d,(710,640),(710,790),"ผ่านด่าน")
    arrow(d,(1390,580),(1440,580),"สถานะศัตรู")
    arrow(d,(900,860),(1030,860),"สถานะเกม")
    arrow(d,(1210,660),(1210,790),"DANGER / ระยะ")
    arrow(d,(1590,660),(1590,790),"intensity")
    arrow(d,(390,590),(1030,845),"ข้อมูลไอเท็ม",label_offset=(20,-25))
    return save(im,"diagram_system_architecture.png")


def player_fsm():
    im,d=canvas("แผนภาพสถานะของผู้เล่น", "PlayerState ใน PlayerController.cs")
    pos={
      "Normal":(650,160,1150,280), "Hidden":(90,390,470,510), "UsingItem":(1330,390,1710,510),
      "Frozen":(90,750,470,870), "Hurt":(710,520,1090,640), "Transitioning":(1330,750,1710,870),
      "Dead":(710,880,1090,1000)
    }
    for key,xy in pos.items():
        thai={"Normal":"ควบคุมปกติ","Hidden":"ซ่อนตัว","UsingItem":"ใช้ไอเท็ม/เปิดกระเป๋า","Frozen":"ถูกหยุดการเคลื่อนไหว","Hurt":"ได้รับความเสียหาย","Transitioning":"กำลังเปลี่ยนฉาก","Dead":"เสียชีวิต"}[key]
        fill=PALE_GREEN if key=="Normal" else PALE_RED if key in ("Hurt","Dead") else PALE
        box(d,xy,key,thai,fill=fill)
    arrow(d,(650,235),(470,420),"กด H / เข้าจุดซ่อน")
    arrow(d,(470,480),(650,255),"กด H / ออกจากจุดซ่อน",label_offset=(-10,18))
    arrow(d,(1150,235),(1330,420),"กด C หรือ Q")
    arrow(d,(1330,480),(1150,255),"ปิดหน้าต่าง",label_offset=(0,18))
    arrow(d,(720,280),(470,750),"ได้รับ Frozen",label_offset=(-18,-8))
    arrow(d,(470,810),(650,260),"ฟื้นตัว",label_offset=(-30,12))
    arrow(d,(900,280),(900,520),"ได้รับความเสียหาย")
    arrow(d,(820,520),(760,280),"พ้นช่วง Hurt",label_offset=(-55,0))
    arrow(d,(1110,260),(1450,750),"ปลดล็อกประตู",label_offset=(30,-10))
    arrow(d,(1330,830),(1110,270),"ยกเลิก/จบด่าน",label_offset=(25,12))
    arrow(d,(900,640),(900,880),"HP เท่ากับ 0")
    arrow(d,(1330,465),(1090,570),"ได้รับความเสียหาย",label_offset=(10,-10))
    d.text((900,1065),"หมายเหตุ: Dead เป็นสถานะสิ้นสุด ส่วน Normal เป็นสถานะหลักที่รับคำสั่งการเล่น",font=font(20),fill=GRAY,anchor="mm")
    return save(im,"diagram_player_fsm.png")


def enemy_fsm():
    im,d=canvas("แผนภาพสถานะของศัตรู", "KillerState ใน KillerAI.cs — แสดงเส้นทางหลักและเหตุการณ์สำคัญ")
    boxes={
      "Patrol":(70,180,330,280), "Alert":(420,180,680,280), "Listening":(770,180,1030,280),
      "InvestigateSound":(1120,180,1510,280), "ConfirmSound":(1120,380,1510,480),
      "Confused":(770,380,1030,480), "Investigate":(420,380,680,480), "Search":(70,380,330,480),
      "Chase":(620,650,900,760), "Attack":(1030,650,1310,760), "AttackRecover":(1410,650,1760,760),
      "ReturnToPatrol":(70,650,480,760), "Scream":(620,900,900,1000)
    }
    for key,xy in boxes.items():
        fill=PALE_GREEN if key=="Patrol" else PALE_RED if key in ("Chase","Attack","Scream") else PALE
        box(d,xy,key,fill=fill,title_size=23)
    arrow(d,(330,230),(420,230),"สงสัย")
    arrow(d,(680,230),(770,230),"ได้ยินเสียง")
    arrow(d,(1030,230),(1120,230),"เลือกเสียง")
    arrow(d,(1315,280),(1315,380),"ถึงตำแหน่ง")
    arrow(d,(1120,430),(1030,430),"ไม่ยืนยัน")
    arrow(d,(770,430),(680,430),"ค้นต่อ")
    arrow(d,(420,430),(330,430),"สูญเสียเป้าหมาย")
    arrow(d,(200,480),(200,650),"หมดเวลาค้นหา")
    arrow(d,(480,705),(620,705),"กลับเส้นลาดตระเวน")
    arrow(d,(900,705),(1030,705),"เข้าเขตโจมตี")
    arrow(d,(1310,705),(1410,705),"โจมตีแล้ว")
    arrow(d,(1585,650),(900,675),"ตรวจพบอีกครั้ง",label_offset=(0,-18))
    arrow(d,(760,760),(760,900),"เข้าเงื่อนไข Scream")
    arrow(d,(620,950),(300,280),"จบ Scream",label_offset=(-15,-5))
    # Global perception shortcuts are dashed to keep the dense FSM readable.
    arrow(d,(200,180),(650,650),"เห็นผู้เล่น",dashed=True,label_offset=(-15,-20),color="#B64B55")
    arrow(d,(1315,480),(860,650),"ยืนยันเสียง/พบเป้าหมาย",dashed=True,label_offset=(15,-18),color="#B64B55")
    d.text((900,1060),"เส้นประสีแดง = เหตุการณ์ตรวจพบผู้เล่นที่สามารถตัดเข้าสู่ Chase จากหลายสถานะ",font=font(20),fill=GRAY,anchor="mm")
    return save(im,"diagram_enemy_fsm.png")


def tension_fsm():
    im,d=canvas("แผนภาพสถานะความตึงเครียด", "TensionState ใน TensionDirector")
    pos={"Calm":(80,460,370,590),"Suspicious":(470,210,850,340),"Hunted":(970,210,1300,340),"Escaped":(1400,460,1720,590),"Cooldown":(690,760,1110,890)}
    for key,xy in pos.items():
        thai={"Calm":"ปกติ","Suspicious":"ศัตรูสงสัย","Hunted":"กำลังถูกไล่ล่า","Escaped":"หลุดจากการไล่ล่า","Cooldown":"ลดระดับความตึงเครียด"}[key]
        fill=PALE_GREEN if key=="Calm" else PALE_RED if key=="Hunted" else PALE
        box(d,xy,key,thai,fill=fill)
    arrow(d,(370,500),(470,310),"เริ่มสงสัย")
    arrow(d,(850,275),(970,275),"เริ่มไล่ล่า")
    arrow(d,(1300,310),(1400,500),"หยุดไล่ล่า")
    arrow(d,(1450,590),(1080,780),"ครบ 1.5 วินาที")
    arrow(d,(690,825),(300,590),"ครบ 2.2 วินาที")
    arrow(d,(780,760),(680,340),"ยังสงสัย")
    arrow(d,(650,340),(820,760),"ไม่สงสัยเกิน 0.8 วินาที",label_offset=(-20,-5))
    arrow(d,(370,540),(970,325),"hunting = true",dashed=True,label_offset=(10,-18),color="#B64B55")
    d.text((900,1020),"สถานะนี้ควบคุมความเข้มของดนตรี เสียงหัวใจ ภาพเตือน และตัวคูณความเร็วศัตรู",font=font(22),fill=GRAY,anchor="mm")
    return save(im,"diagram_tension_fsm.png")


def insert_diagrams(paths):
    doc=Document(INPUT)
    replacements={
      "[แทรกภาพ: System Architecture Diagram]":paths[0],
      "[แทรกภาพ: Player FSM Diagram]":paths[1],
      "[แทรกภาพ: Enemy FSM Diagram]":paths[2],
      "[แทรกภาพ: Tension State Diagram]":paths[3],
    }
    found=[]
    for p in doc.paragraphs:
        key=p.text.strip()
        if key in replacements:
            p.clear()
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Inches(0.08)
            p.paragraph_format.space_after=Inches(0.08)
            run=p.add_run()
            run.add_picture(str(replacements[key]),width=Inches(6.0))
            found.append(key)
    missing=set(replacements)-set(found)
    if missing:
        raise RuntimeError(f"Missing placeholders: {sorted(missing)}")
    doc.core_properties.title="บทที่ 3 วิธีการดำเนินงาน ฉบับปรับปรุงพร้อมไดอะแกรม"
    doc.save(OUTPUT)
    print(OUTPUT)
    print("inserted",len(found))


if __name__=="__main__":
    paths=[architecture(),player_fsm(),enemy_fsm(),tension_fsm()]
    insert_diagrams(paths)
